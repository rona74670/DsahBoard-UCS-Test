"""
FastAPI backend for UCS Manager Dashboard  v4.0
Multi-farm: TC-FARM · PT-FARM · FATTAL
Database: SQLite (dashboard.db)
"""

import re
import collections
import io
import threading
import os
import shutil
import time
import secrets
import urllib.parse
from datetime import datetime, timedelta
from typing import Optional
from concurrent.futures import ThreadPoolExecutor, as_completed
from fastapi import FastAPI, HTTPException, APIRouter, UploadFile, File, Depends, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, JSONResponse, RedirectResponse
from pydantic import BaseModel
from dotenv import load_dotenv
import openpyxl
import requests

# Load .env before any local module that reads env vars at import time
load_dotenv(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env"))

from ucs_client import UCSMClient
from infinidat_client import InfinidatClient
from vcenter_client import VCenterClient
from veeam_client import VeeamEMClient
from veeam_br_client import VeeamBRClient
import ollama_client
from database import (
    init_db, migrate_json_if_needed,
    load_manual, save_manual,
    get_vm_notes, save_vm_note,
    get_vm_classifications, save_vm_classification,
    list_devices, get_device, create_device, update_device, delete_device,
    record_ibox_capacity, get_ibox_capacity_history, compute_days_until_full,
    list_customers, get_customers_map, get_customer,
    create_customer, update_customer, delete_customer, upsert_customers_bulk,
    insert_vcenter_events, get_vcenter_events, count_vcenter_events_by_env,
    get_microsoft_auth_config, save_microsoft_auth_config,
    list_authz_users, upsert_authz_user, delete_authz_user,
    resolve_authz_user, authz_counts,
    create_vm_snapshot, insert_snapshot_items, finalize_snapshot,
    insert_snapshot_env_data, get_snapshot_env_data,
    list_snapshots, get_snapshot, get_snapshot_items, delete_snapshot,
    prune_old_snapshots, has_protected_snapshot_today, has_snapshot_today_by,
    list_billing_configs, get_billing_config, upsert_billing_config, delete_billing_config,
)
from auth import require_admin, get_current_user, create_access_token, decode_token, JWT_EXPIRE_H

app = FastAPI(title="UCS Dashboard API", version="4.0.0")

# Auth middleware (install_auth.ps1)
from starlette.middleware.base import BaseHTTPMiddleware
from starlette.requests import Request as _SReq
from starlette.responses import JSONResponse as _SResp

class _AuthMiddleware(BaseHTTPMiddleware):
    _API_OPEN = frozenset({
        "/api/auth/login",
        "/api/auth/logout",
        "/api/auth/login-options",
        "/api/auth/microsoft/start",
        "/api/auth/microsoft/callback",
    })
    # "/health" doesn't start with "/api/", so it's checked against this set,
    # not _API_OPEN, and must return 200 for unauthenticated monitoring probes.
    _PAGE_OPEN = frozenset({"/login", "/health"})

    async def dispatch(self, request: _SReq, call_next):
        path = request.url.path

        token = request.cookies.get("auth_token", "")
        if not token:
            auth = request.headers.get("authorization", "")
            token = auth[7:] if auth.lower().startswith("bearer ") else ""

        is_valid = False
        if token:
            try:
                payload = decode_token(token)
                is_valid = bool(payload.get("sub"))
            except Exception:
                is_valid = False

        if path.startswith("/api/") and path not in self._API_OPEN and not is_valid:
            return _SResp(
                {"detail": "Authentication required"},
                status_code=401,
                headers={"WWW-Authenticate": "Bearer"},
            )

        if (not path.startswith("/api/") and
            not path.startswith("/static/") and
            path not in self._PAGE_OPEN and
            not is_valid):
            target = "/login?next=" + urllib.parse.quote(path)
            return RedirectResponse(url=target, status_code=302)

        return await call_next(request)

app.add_middleware(_AuthMiddleware)


def _safe_next_path(value) -> str:
    """Only allow same-origin relative redirect targets. Rejects protocol-relative
    ("//evil.com") and backslash ("/\\evil.com") tricks that some browsers resolve
    as scheme-relative URLs, which would otherwise send users to an attacker site
    right after login."""
    if not isinstance(value, str) or not value:
        return "/"
    value = value.strip()
    if not value.startswith("/") or value.startswith(("//", "/\\")):
        return "/"
    return value

_CORS_ORIGINS = [o.strip() for o in os.getenv("CORS_ALLOWED_ORIGINS", "").split(",") if o.strip()]
if not _CORS_ORIGINS:
    # Override via CORS_ALLOWED_ORIGINS env var for multi-hostname setups
    _CORS_ORIGINS = ["https://localhost:9000", "https://tatayas.allegronet.co.il:9000"]


class _CSRFMiddleware(BaseHTTPMiddleware):
    """Double-submit-cookie CSRF protection for mutating API endpoints.
    Exempts Bearer-token requests (not sent automatically by the browser,
    so not CSRF-able) for backward compatibility during migration."""
    _MUTATING = frozenset({"POST", "PUT", "DELETE", "PATCH"})
    _EXEMPT   = frozenset({"/api/auth/login", "/api/auth/logout"})

    async def dispatch(self, request: _SReq, call_next):
        path = request.url.path
        if (request.method in self._MUTATING and
                path.startswith("/api/") and
                path not in self._EXEMPT and
                request.cookies.get("auth_token")):
            cookie_csrf = request.cookies.get("csrf_token", "")
            header_csrf = request.headers.get("x-csrf-token", "")
            if not cookie_csrf or not header_csrf or not secrets.compare_digest(cookie_csrf, header_csrf):
                return _SResp({"detail": "CSRF token missing or invalid"}, status_code=403)
        return await call_next(request)


class _SecurityHeadersMiddleware(BaseHTTPMiddleware):
    """Add security headers to every response."""
    async def dispatch(self, request: _SReq, call_next):
        response = await call_next(request)
        response.headers["X-Content-Type-Options"] = "nosniff"
        response.headers["X-Frame-Options"] = "DENY"
        response.headers["X-XSS-Protection"] = "1; mode=block"
        response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
        response.headers["Content-Security-Policy"] = (
            "default-src 'self'; "
            "script-src 'self' 'unsafe-inline'; "
            "style-src 'self' 'unsafe-inline'; "
            "img-src 'self' data: blob:; "
            "connect-src 'self'; "
            "font-src 'self' data:; "
            "object-src 'none'; "
            "frame-ancestors 'none';"
        )
        return response

app.add_middleware(_SecurityHeadersMiddleware)
app.add_middleware(_CSRFMiddleware)

app.add_middleware(
    CORSMiddleware,
    allow_origins=_CORS_ORIGINS,
    allow_credentials=True,
    allow_methods=["GET", "POST", "PUT", "DELETE"],
    allow_headers=["Authorization", "Content-Type"],
)

BACKEND_DIR  = os.path.dirname(os.path.abspath(__file__))
FRONTEND_DIR = os.path.normpath(os.path.join(BACKEND_DIR, "..", "frontend"))

init_db()


# ── Login rate limiting ──────────────────────────────────────────────────────
_LOGIN_ATTEMPTS: dict = {}
_LOGIN_WINDOW    = 60   # seconds
_LOGIN_MAX_TRIES = 10   # per window per IP


def _check_rate_limit(ip: str):
    now = time.time()
    q   = _LOGIN_ATTEMPTS.setdefault(ip, collections.deque())
    while q and now - q[0] > _LOGIN_WINDOW:
        q.popleft()
    if len(q) >= _LOGIN_MAX_TRIES:
        print(f"[auth] rate-limit exceeded: {ip}")
        raise HTTPException(status_code=429, detail="Too many login attempts. Try again later.")
    q.append(now)


# Trusted reverse proxies — only their X-Forwarded-For header is accepted
_TRUSTED_PROXIES: frozenset = frozenset(
    ip.strip() for ip in os.getenv("TRUSTED_PROXIES", "127.0.0.1,::1").split(",") if ip.strip()
)
_LOCAL_ONLY_USERS: frozenset = frozenset(
    u.strip().lower() for u in os.getenv("LOCAL_ONLY_USERS", "rona").split(",") if u.strip()
)
_FORCED_ADMIN_IDENTITIES: frozenset = frozenset(
    u.strip().lower() for u in os.getenv(
        "FORCED_ADMIN_IDENTITIES",
        "rona@allegronet.co.il,rona",
    ).split(",") if u.strip()
)
_LOCAL_LOGIN_SOURCE_IPS: frozenset = frozenset(
    ip.strip() for ip in os.getenv(
        "LOCAL_LOGIN_SOURCE_IPS",
        "127.0.0.1,::1,10.114.6.11",
    ).split(",") if ip.strip()
)
_PERMISSIONS_MANAGERS: frozenset = frozenset(
    u.strip().lower() for u in os.getenv(
        "PERMISSIONS_MANAGERS",
        "rona@allegronet.co.il,rona",
    ).split(",") if u.strip()
)


def _is_loopback_ip(ip: str) -> bool:
    return ip in {"127.0.0.1", "::1", "localhost"}


def _normalize_identity(identity: str) -> str:
    s = str(identity or "").strip().lower()
    if "\\" in s:
        s = s.split("\\", 1)[1].strip()
    return s


def _is_forced_admin(identity: str) -> bool:
    ident = _normalize_identity(identity)
    if not ident:
        return False
    if ident in _FORCED_ADMIN_IDENTITIES:
        return True
    short = ident.split("@", 1)[0].strip()
    return bool(short) and short in _FORCED_ADMIN_IDENTITIES


def _is_permissions_manager(identity: str) -> bool:
    ident = _normalize_identity(identity)
    if not ident:
        return False
    if ident in _PERMISSIONS_MANAGERS:
        return True
    if _is_forced_admin(ident):
        return True

    short = ident.split("@", 1)[0].strip()
    if short and short in _PERMISSIONS_MANAGERS:
        return True

    # Local-only users are allowed to manage permissions as requested.
    return bool(short) and short in _LOCAL_ONLY_USERS


def require_permissions_manager(user: dict = Depends(get_current_user)) -> str:
    username = str(user.get("username", ""))
    if not _is_permissions_manager(username):
        raise HTTPException(status_code=403, detail="Permissions management is restricted")
    return username


def _is_local_login_source_ip(ip: str) -> bool:
    ip_s = str(ip or "").strip()
    return ip_s in _LOCAL_LOGIN_SOURCE_IPS or _is_loopback_ip(ip_s)


def _get_client_ip(request) -> str:
    """Return the real client IP, honouring X-Forwarded-For only from trusted proxies."""
    real_ip = request.client.host if request.client else "unknown"
    if real_ip in _TRUSTED_PROXIES:
        fwd = request.headers.get("x-forwarded-for", "")
        return fwd.split(",")[0].strip() if fwd else real_ip
    return real_ip


def _check_admin_refresh(refresh: bool, request) -> None:
    """Raise 403 when refresh=True and the caller is not admin."""
    if not refresh:
        return
    token = request.cookies.get("auth_token", "")
    if not token:
        auth = request.headers.get("authorization", "")
        token = auth[7:] if auth.lower().startswith("bearer ") else ""
    if token:
        try:
            from auth import JWT_SECRET, JWT_ALGO
            from jose import jwt as _j
            payload = _j.decode(token, JWT_SECRET, algorithms=[JWT_ALGO])
            if payload.get("role") == "admin":
                return
        except Exception:
            pass
    raise HTTPException(status_code=403, detail="Forced cache refresh requires admin role")


class LoginRequest(BaseModel):
    username: str
    password: str


class AuthzUserIn(BaseModel):
    identity: str
    role: str = "readonly"


def _normalize_role(role: str) -> str:
    r = str(role or "").strip().lower()
    return "admin" if r == "admin" else "readonly"


def _is_local_request(request: Request) -> bool:
    return _is_local_login_source_ip(_get_client_ip(request))


@app.get("/api/auth/login-options")
def api_login_options(request: Request):
    cfg = get_microsoft_auth_config()
    is_local = _is_local_request(request)
    return {
        "is_local": is_local,
        "password_login_enabled": is_local,
        "entra_enabled": bool(cfg.get("enabled")),
    }


@app.post("/api/auth/login")
def api_login(req: LoginRequest, request: Request):
    from auth import authenticate_ad, authenticate_local, JWT_ALGO, JWT_SECRET
    client_ip = _get_client_ip(request)
    if not _is_local_request(request):
        raise HTTPException(
            status_code=403,
            detail="Username/password login is local-only. Use Entra ID sign-in.",
        )
    _check_rate_limit(client_ip)
    result = authenticate_ad(req.username, req.password)
    local_result = None
    if not result:
        local_result = authenticate_local(req.username, req.password)
        result = local_result

    if not result:
        print(f"[auth] failed login: user='{req.username}' ip={client_ip}")
        raise HTTPException(status_code=401, detail="Invalid username or password")

    if local_result and local_result.get("username", "").lower() in _LOCAL_ONLY_USERS:
        real_ip = _get_client_ip(request)
        if not _is_local_login_source_ip(real_ip):
            raise HTTPException(
                status_code=403,
                detail="Local account login is allowed only from the server itself",
            )

    if _is_forced_admin(result.get("username", "")):
        result["role"] = "admin"

    token = create_access_token(result["username"], result["role"])
    # Extract expiry from the token for the client-side session timer
    try:
        from jose import jwt as _jt
        _exp = _jt.decode(token, JWT_SECRET, algorithms=[JWT_ALGO]).get("exp", 0)
    except Exception:
        _exp = 0
    _csrf = secrets.token_hex(32)
    resp  = JSONResponse({
        "username":   result["username"],
        "role":       result["role"],
        "expires_at": _exp,           # Unix timestamp — no raw JWT returned
    })
    resp.set_cookie(
        key="auth_token", value=token,
        httponly=True, secure=True, samesite="lax",
        max_age=JWT_EXPIRE_H * 3600, path="/",
    )
    resp.set_cookie(
        key="csrf_token", value=_csrf,
        httponly=False, secure=True, samesite="lax",
        max_age=JWT_EXPIRE_H * 3600, path="/",
    )
    return resp


@app.post("/api/auth/logout")
def api_logout():
    resp = JSONResponse({"status": "logged out"})
    resp.delete_cookie(key="auth_token", path="/")
    resp.delete_cookie(key="csrf_token", path="/")
    return resp


@app.get("/api/auth/session")
def api_auth_session(request: Request):
    token = request.cookies.get("auth_token", "")
    if not token:
        raise HTTPException(status_code=401, detail="Not authenticated")
    try:
        payload = decode_token(token)
    except Exception:
        raise HTTPException(status_code=401, detail="Invalid session")
    return {
        "username": payload.get("sub", ""),
        "role": payload.get("role", "user"),
        "expires_at": payload.get("exp", 0),
    }


class MicrosoftSsoConfigIn(BaseModel):
    enabled: bool = False
    tenant_id: str = ""
    client_id: str = ""
    client_secret: str = ""
    redirect_uri: str = ""
    scope: str = "openid profile email User.Read"
    ssl_enabled: bool = True
    ssl_mode: str = "https_only"


def _microsoft_cfg_masked() -> dict:
    cfg = get_microsoft_auth_config()
    secret = cfg.get("client_secret", "")
    cfg["client_secret"] = "" if not secret else "********"
    cfg["has_client_secret"] = bool(secret)
    return cfg


@app.get("/api/auth/microsoft/config")
def api_microsoft_config_get(_: str = Depends(require_admin)):
    return _microsoft_cfg_masked()


@app.put("/api/auth/microsoft/config")
def api_microsoft_config_put(payload: MicrosoftSsoConfigIn, _: str = Depends(require_admin)):
    cfg = payload.model_dump()
    if cfg.get("enabled"):
        for field in ("tenant_id", "client_id"):
            if not str(cfg.get(field, "")).strip():
                raise HTTPException(status_code=400, detail=f"Missing required field: {field}")

    existing = get_microsoft_auth_config()
    # Keep previous secret when the admin leaves it empty.
    if not str(cfg.get("client_secret", "")).strip():
        cfg["client_secret"] = existing.get("client_secret", "")

    save_microsoft_auth_config(cfg)
    return {"status": "saved", "config": _microsoft_cfg_masked()}


@app.get("/api/auth/microsoft/start")
def api_microsoft_start(request: Request, next: str = "/"):
    cfg = get_microsoft_auth_config()
    if not cfg.get("enabled"):
        raise HTTPException(status_code=400, detail="Microsoft SSO is disabled")

    tenant_id = str(cfg.get("tenant_id", "")).strip()
    client_id = str(cfg.get("client_id", "")).strip()
    client_secret = str(cfg.get("client_secret", "")).strip()
    if not tenant_id or not client_id or not client_secret:
        raise HTTPException(status_code=400, detail="Microsoft SSO configuration is incomplete")

    redirect_uri = str(cfg.get("redirect_uri", "")).strip()
    if not redirect_uri:
        redirect_uri = str(request.url_for("api_microsoft_callback"))

    state = secrets.token_urlsafe(32)
    scope = str(cfg.get("scope", "openid profile email User.Read")).strip() or "openid profile email User.Read"
    q = {
        "client_id": client_id,
        "response_type": "code",
        "redirect_uri": redirect_uri,
        "response_mode": "query",
        "scope": scope,
        "state": state,
    }
    auth_url = (
        f"https://login.microsoftonline.com/{tenant_id}/oauth2/v2.0/authorize?" +
        urllib.parse.urlencode(q)
    )
    resp = RedirectResponse(url=auth_url, status_code=302)
    resp.set_cookie(
        key="ms_state",
        value=state,
        httponly=True,
        secure=True,
        samesite="lax",
        max_age=600,
        path="/",
    )
    safe_next = _safe_next_path(next)
    resp.set_cookie(
        key="ms_next",
        value=safe_next,
        httponly=True,
        secure=True,
        samesite="lax",
        max_age=600,
        path="/",
    )
    return resp


@app.get("/api/auth/microsoft/callback")
def api_microsoft_callback(request: Request, code: str = "", state: str = ""):
    expected_state = request.cookies.get("ms_state", "")
    if not code or not state or state != expected_state:
        raise HTTPException(status_code=400, detail="Invalid Microsoft callback state")

    cfg = get_microsoft_auth_config()
    tenant_id = str(cfg.get("tenant_id", "")).strip()
    client_id = str(cfg.get("client_id", "")).strip()
    client_secret = str(cfg.get("client_secret", "")).strip()
    redirect_uri = str(cfg.get("redirect_uri", "")).strip() or str(request.url_for("api_microsoft_callback"))
    scope = str(cfg.get("scope", "openid profile email User.Read")).strip() or "openid profile email User.Read"

    token_url = f"https://login.microsoftonline.com/{tenant_id}/oauth2/v2.0/token"
    tr = requests.post(
        token_url,
        data={
            "client_id": client_id,
            "client_secret": client_secret,
            "code": code,
            "redirect_uri": redirect_uri,
            "grant_type": "authorization_code",
            "scope": scope,
        },
        timeout=30,
    )
    if tr.status_code >= 400:
        raise HTTPException(status_code=502, detail="Microsoft token exchange failed")

    body = tr.json()
    id_token = body.get("id_token", "")
    if not id_token:
        raise HTTPException(status_code=502, detail="Missing ID token from Microsoft")

    try:
        from jose import jwt as _jwt
        claims = _jwt.get_unverified_claims(id_token)
    except Exception:
        claims = {}

    username = (
        claims.get("preferred_username") or
        claims.get("upn") or
        claims.get("email") or
        claims.get("oid") or
        "entra-user"
    )

    if _is_forced_admin(str(username)):
        role = "admin"
    else:
        authz_user = resolve_authz_user(str(username))
        if not authz_user:
            return RedirectResponse(url="/login?error=unauthorized_user", status_code=302)

        role = _normalize_role(authz_user.get("role", "readonly"))

    token = create_access_token(str(username), role)
    exp = 0
    try:
        from jose import jwt as _jt
        from auth import JWT_ALGO, JWT_SECRET
        exp = _jt.decode(token, JWT_SECRET, algorithms=[JWT_ALGO]).get("exp", 0)
    except Exception:
        pass

    next_url = _safe_next_path(request.cookies.get("ms_next", "/"))
    resp = RedirectResponse(url=next_url, status_code=302)
    resp.set_cookie(
        key="auth_token",
        value=token,
        httponly=True,
        secure=True,
        samesite="lax",
        max_age=JWT_EXPIRE_H * 3600,
        path="/",
    )
    resp.set_cookie(
        key="csrf_token",
        value=secrets.token_hex(32),
        httponly=False,
        secure=True,
        samesite="lax",
        max_age=JWT_EXPIRE_H * 3600,
        path="/",
    )
    resp.delete_cookie(key="ms_state", path="/")
    resp.delete_cookie(key="ms_next", path="/")
    resp.set_cookie(
        key="login_meta",
        value=str(exp),
        httponly=False,
        secure=True,
        samesite="lax",
        max_age=JWT_EXPIRE_H * 3600,
        path="/",
    )
    return resp


@app.get("/api/authz/users")
def api_authz_users(_: str = Depends(require_permissions_manager)):
    return {
        "policy": "allow_list_only",
        "users": list_authz_users(),
        "counts": authz_counts(),
    }


@app.put("/api/authz/users")
def api_authz_users_upsert(payload: AuthzUserIn, _: str = Depends(require_permissions_manager)):
    identity = str(payload.identity or "").strip()
    if not identity:
        raise HTTPException(status_code=400, detail="identity is required")
    role = _normalize_role(payload.role)
    try:
        row = upsert_authz_user(identity=identity, role=role)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"status": "saved", "user": row, "counts": authz_counts()}


@app.delete("/api/authz/users/{identity}")
def api_authz_users_delete(identity: str, _: str = Depends(require_permissions_manager)):
    ident = urllib.parse.unquote(identity)
    row = resolve_authz_user(ident)
    if not row:
        raise HTTPException(status_code=404, detail="User not found")

    counts = authz_counts()
    if _normalize_role(row.get("role", "readonly")) == "admin" and counts.get("admin", 0) <= 1:
        raise HTTPException(status_code=400, detail="At least one edit user must remain")

    ok = delete_authz_user(ident)
    if not ok:
        raise HTTPException(status_code=404, detail="User not found")
    return {"status": "deleted", "counts": authz_counts()}



# ------------------------------------------------------------------ #
#  Pydantic models                                                     #
# ------------------------------------------------------------------ #

class ManualEntry(BaseModel):
    server_key: str
    assignment: str = ""
    cabinet:    str = ""
    vcenter:    str = ""
    notes:      str = ""


class DeviceIn(BaseModel):
    farm:        str  = "ALL"
    name:        str
    type:        str  = "Generic"
    role:        str  = ""
    host:        str  = ""
    ip_address:  str  = ""
    ilo_address: str  = ""
    username:    str  = ""
    password:    str  = ""
    port:        int  = 443
    protocol:    str  = "https"
    notes:       str  = ""
    enabled:     bool = True




class CustomerIn(BaseModel):
    sap_number:    str
    customer_name: str


class VmNoteIn(BaseModel):
    vm_name: str
    note:    str = ""


class AiChatMessage(BaseModel):
    role:    str
    content: str


class AiChatIn(BaseModel):
    message: str
    history: list[AiChatMessage] = []
    context: str = ""


VM_CLASSIFICATIONS = [
    "שרת אלגרונט בלבד",
    "שרת בשימוש אלגרונט + לקוחות",
    "שרת לקוחות תחת דומיין אלגרונט",
    "שרת אלגרונט אבטחת מידע",
    "שרת אלגרונט ניטור",
    "שרת אלגרונט גיבוי",
    "שרת לוגים",
    "שרת דומיין פרטי",
    "מרכזיה",
    "שרת לניהול סטוראג אלגרונט",
    "שרת VMware אלגרונט",
    "אלגרונט החדשה",
    "DR של תל אביב",
    "אופטימה ישרוטל",
    "הילטון",
    "יוון",
]


class VmClassificationIn(BaseModel):
    vm_name:        str
    classification: str = ""


# ------------------------------------------------------------------ #
#  Farm router factory                                                 #
# ------------------------------------------------------------------ #

def build_farm_router(prefix: str, host: str, user: str, pwd: str, farm_id: str) -> APIRouter:
    router = APIRouter(prefix=prefix)

    def _client():
        return UCSMClient(host=host, username=user, password=pwd)

    def _run(func):
        client = _client()
        try:
            client.login()
            return func(client)
        except RuntimeError as e:
            raise HTTPException(status_code=502, detail="Backend communication error")
        except Exception as e:
            raise HTTPException(status_code=500, detail="Unexpected error")
        finally:
            client.logout()

    @router.get("/ping")
    def ping():
        client = _client()
        try:
            client.login()
            client.logout()
            return {"status": "ok", "host": host, "user": user}
        except RuntimeError as e:
            raise HTTPException(status_code=502, detail="Backend communication error")
        except Exception as e:
            raise HTTPException(status_code=500, detail="Connection failed")

    @router.get("/summary")
    def get_summary():
        return _run(lambda c: c.get_summary())

    @router.get("/system")
    def get_system():
        return _run(lambda c: c.get_system_info())

    @router.get("/fabric-interconnects")
    def get_fis():
        return _run(lambda c: c.get_fabric_interconnects())

    @router.get("/chassis")
    def get_chassis():
        return _run(lambda c: c.get_chassis())

    @router.get("/blades")
    def get_blades():
        return _run(lambda c: c.get_blade_servers())

    @router.get("/racks")
    def get_racks():
        return _run(lambda c: c.get_rack_servers())

    @router.get("/service-profiles")
    def get_sps():
        return _run(lambda c: c.get_service_profiles())

    @router.get("/faults")
    def get_faults(min_severity: str = "warning"):
        return _run(lambda c: c.get_faults(min_severity=min_severity))

    @router.get("/blades/enriched")
    def get_blades_enriched():
        return _run(lambda c: c.get_blades_enriched())

    @router.get("/hardware-summary")
    def get_hardware_summary():
        return _run(lambda c: c.get_hardware_summary())

    @router.get("/firmware")
    def get_firmware():
        return _run(lambda c: c.get_firmware_summary())

    @router.get("/power-stats")
    def get_power_stats():
        return _run(lambda c: c.get_power_stats())

    @router.get("/network-adapters")
    def get_network_adapters():
        return _run(lambda c: c.get_network_summary())

    @router.get("/pools")
    def get_pools():
        return _run(lambda c: c.get_pool_data())

    @router.get("/manual-data")
    def get_manual_data():
        return load_manual(farm_id)

    @router.put("/manual-data")
    def update_manual_data(entry: ManualEntry):
        save_manual(
            farm_id, entry.server_key,
            assignment=entry.assignment,
            cabinet=entry.cabinet,
            vcenter=entry.vcenter,
            notes=entry.notes,
        )
        return {"status": "saved", "key": entry.server_key}

    return router


# ------------------------------------------------------------------ #
#  Register UCS farms                                                  #
# ------------------------------------------------------------------ #

TC_HOST   = os.getenv("TC_HOST",     "10.10.140.100")
TC_USER   = os.getenv("TC_USER",     os.getenv("UCS_USER", ""))
TC_PASS   = os.getenv("TC_PASS",     os.getenv("UCS_PASS", ""))

PT_HOST   = os.getenv("PT_HOST",     "10.20.140.100")
PT_USER   = os.getenv("PT_USER",     os.getenv("UCS_USER", ""))
PT_PASS   = os.getenv("PT_PASS",     os.getenv("UCS_PASS", ""))

FATTAL_HOST   = os.getenv("FATTAL_HOST", "10.8.0.160")
FATTAL_USER   = os.getenv("FATTAL_USER", os.getenv("UCS_USER", ""))
FATTAL_PASS   = os.getenv("FATTAL_PASS", os.getenv("UCS_PASS", ""))

migrate_json_if_needed("tc",     os.path.join(BACKEND_DIR, "manual_data_tc.json"))
migrate_json_if_needed("pt",     os.path.join(BACKEND_DIR, "manual_data_pt.json"))
migrate_json_if_needed("fattal", os.path.join(BACKEND_DIR, "manual_data_fattal.json"))
_old = os.path.join(BACKEND_DIR, "manual_data.json")
if os.path.exists(_old):
    migrate_json_if_needed("tc", _old)

app.include_router(build_farm_router("/api/tc",     TC_HOST,     TC_USER,     TC_PASS,     "tc"))
app.include_router(build_farm_router("/api/pt",     PT_HOST,     PT_USER,     PT_PASS,     "pt"))
app.include_router(build_farm_router("/api/fattal", FATTAL_HOST, FATTAL_USER, FATTAL_PASS, "fattal"))

# ── Infinidat iBox credentials ──────────────────────────────────────────
IBOX_TC03_HOST = os.getenv("IBOX_TC03_HOST", "ibox-tc-03.allegronet.co.il")
IBOX_TC03_USER = os.getenv("IBOX_TC03_USER", os.getenv("UCS_USER", ""))
IBOX_TC03_PASS = os.getenv("IBOX_TC03_PASS", os.getenv("UCS_PASS", ""))

IBOX_TC_OLD_HOST = os.getenv("IBOX_TC_OLD_HOST", "10.10.15.10")
IBOX_TC_OLD_USER = os.getenv("IBOX_TC_OLD_USER", os.getenv("UCS_USER", ""))
IBOX_TC_OLD_PASS = os.getenv("IBOX_TC_OLD_PASS", os.getenv("UCS_PASS", ""))

IBOX_PT_HOST = os.getenv("IBOX_PT_HOST", "10.10.25.200")
IBOX_PT_USER = os.getenv("IBOX_PT_USER", os.getenv("UCS_USER", ""))
IBOX_PT_PASS = os.getenv("IBOX_PT_PASS", os.getenv("UCS_PASS", ""))

def build_ibox_router(prefix: str, host: str, user: str, pwd: str, ibox_id: str) -> APIRouter:
    router = APIRouter(prefix=prefix)
    client = InfinidatClient(host=host, username=user, password=pwd)

    def _run(func):
        try:
            return func(client)
        except Exception as e:
            print(f"[vcenter:{prefix}] {type(e).__name__}: {e}")
            raise HTTPException(status_code=502, detail="Backend communication error") from e

    @router.get("/ping")
    def ping():
        try:
            sys = client.get_system()
            return {"status": "ok", "host": host, "version": sys.get("version", "")}
        except Exception as e:
            raise HTTPException(status_code=502, detail="Backend communication error")

    @router.get("/probe-volume-metrics", dependencies=[Depends(require_admin)])
    def probe_volume_metrics():
        """Read-only probe: does this iBox expose per-volume latency, and what does a volume's serial look like?"""
        import json as _json
        out = {"host": host}
        try:
            out["version"] = client.get_system().get("version", "")
            vols = client._get("volumes", {"type": "MASTER", "page_size": 3}).get("result") or []
        except Exception as e:
            raise HTTPException(status_code=502, detail="Backend communication error")
        out["sample_volumes"] = [{k: v.get(k) for k in ("id", "name", "serial", "pool_name", "size", "type")} for v in vols]
        out["volume_keys"] = sorted(vols[0].keys()) if vols else []
        if not vols:
            return out
        vid = vols[0]["id"]
        probes = {}
        for path in (f"volumes/{vid}/metrics", f"volumes/{vid}/performance", f"volumes/{vid}/stats",
                     "metrics", "metrics/available", "volumes/performance", "volumes?fields=id,name,serial,performance"):
            try:
                r = client._session.get(f"{client.base}/{path}", timeout=20)
                try:
                    body = _json.dumps(r.json(), ensure_ascii=False)[:600]
                except Exception:
                    body = r.text[:300]
                probes[path] = {"status": r.status_code, "sample": body}
            except Exception as e:
                probes[path] = {"status": "error", "sample": type(e).__name__}
        out["probes"] = probes
        return out

    @router.get("/summary")
    def summary():
        data = _run(lambda c: c.get_summary())
        try:
            cap = data.get("capacity", {})
            record_ibox_capacity(
                ibox_id,
                phys_free_tb=cap.get("phys_free_tb", 0),
                phys_total_tb=cap.get("phys_total_tb", 0),
                phys_pct=cap.get("phys_pct", 0),
            )
        except Exception:
            pass
        return data

    @router.get("/pools")
    def pools():
        return _run(lambda c: c.get_pools())

    @router.get("/volumes")
    def volumes():
        return _run(lambda c: c.get_volumes())

    @router.get("/snapshots")
    def snapshots():
        return _run(lambda c: c.get_snapshots())

    @router.get("/health")
    def health():
        return _run(lambda c: c.get_health())

    @router.get("/replicas")
    def replicas():
        return _run(lambda c: c.get_replicas())

    @router.get("/replicas/detailed")
    def replicas_detailed():
        return _run(lambda c: c.get_replicas_detailed())

    @router.get("/events")
    def events():
        return _run(lambda c: c.get_events())

    @router.get("/top-volumes")
    def top_volumes():
        return _run(lambda c: c.get_top_volumes())

    @router.get("/cgs")
    def cgs():
        return _run(lambda c: c.get_cgs())

    @router.get("/hosts")
    def hosts():
        return _run(lambda c: c.get_hosts())

    @router.get("/capacity-trend")
    def capacity_trend():
        history = get_ibox_capacity_history(ibox_id, days=30)
        days_left = compute_days_until_full(history)
        return {
            "ibox_id": ibox_id,
            "readings": len(history),
            "days_until_full": days_left,
            "history": history,
        }

    @router.get("/performance")
    def performance():
        return _run(lambda c: c.get_performance())

    @router.get("/hardware")
    def hardware():
        return _run(lambda c: c.get_hardware())

    return router

app.include_router(build_ibox_router("/api/ibox/tc03",  IBOX_TC03_HOST,  IBOX_TC03_USER,  IBOX_TC03_PASS,  "tc03"))
app.include_router(build_ibox_router("/api/ibox/tc-old", IBOX_TC_OLD_HOST, IBOX_TC_OLD_USER, IBOX_TC_OLD_PASS, "tc-old"))
app.include_router(build_ibox_router("/api/ibox/pt",     IBOX_PT_HOST,     IBOX_PT_USER,     IBOX_PT_PASS,     "pt"))


# ------------------------------------------------------------------ #
#  vCenter credentials                                                 #
# ------------------------------------------------------------------ #

def _clean_env(v: str | None) -> str:
    if v is None:
        return ""
    s = v.strip()
    if len(s) >= 2 and s[0] == s[-1] and s[0] in ("\"", "'"):
        s = s[1:-1].strip()
    return s


def _env_first(*keys: str, default: str = "") -> str:
    for k in keys:
        val = os.getenv(k)
        if val is not None and val.strip() != "":
            return _clean_env(val)
    return default


VCENTER_TC_HOST = _env_first("VCENTER_TC_HOST", "VCENTER_TC")
VCENTER_TC_USER = _env_first("VCENTER_TC_USER", "VCENTER_TC_USERNAME")
VCENTER_TC_PASS = _env_first("VCENTER_TC_PASS", "VCENTER_TC_PASSWORD")

VCENTER_PT_HOST = _env_first("VCENTER_PT_HOST", "VCENTER_PT", default="al-vcenter6.allegronet.co.il")
VCENTER_PT_USER = _env_first("VCENTER_PT_USER", "VCENTER_PT_USERNAME")
VCENTER_PT_PASS = _env_first("VCENTER_PT_PASS", "VCENTER_PT_PASSWORD")

VCENTER_TC8_HOST = _env_first("VCENTER_TC8_HOST", "VCENTER_TC8", default="vcenter-tc.allegronet.co.il")
VCENTER_TC8_USER = _env_first("VCENTER_TC8_USER", "VCENTER_TC8_USERNAME", default=r"VSPHERE.TC\CockpitDash")
VCENTER_TC8_PASS = _env_first("VCENTER_TC8_PASS", "VCENTER_TC8_PASSWORD", default=VCENTER_TC_PASS)

VCENTER_FATTAL_HOST = _env_first("VCENTER_FATTAL_HOST", "VCENTER_FATTAL", default="")
VCENTER_FATTAL_USER = _env_first("VCENTER_FATTAL_USER", "VCENTER_FATTAL_USERNAME", "VSPHERE_FATTAL_USER")
VCENTER_FATTAL_PASS = _env_first("VCENTER_FATTAL_PASS", "VCENTER_FATTAL_PASSWORD", "VSPHERE_FATTAL_PASS")

_VCENTER_CACHE_TTL = 180

def build_vcenter_router(prefix: str, host: str, user: str, pwd: str) -> APIRouter:
    router = APIRouter(prefix=prefix)
    env_key = prefix.rsplit("/", 1)[-1]
    _cache: dict = {}
    _cache_lock = threading.Lock()

    def _client():
        if not host:
            raise HTTPException(status_code=503, detail="vCenter host not configured")
        return VCenterClient(host=host, username=user, password=pwd)

    def _run(func):
        client = _client()
        try:
            return func(client)
        except HTTPException:
            raise
        except Exception as e:
            print(f"[vcenter:{env_key}] {type(e).__name__}: {e}")
            raise HTTPException(status_code=502, detail=f"Backend communication error: {e}")

    @router.get("/ping")
    def ping():
        return _run(lambda c: c.ping())

    @router.get("/vms")
    def vms():
        return _run(lambda c: c.get_vms())

    @router.get("/summary")
    def summary(refresh: bool = False, request: Request = None):
        _check_admin_refresh(refresh, request)
        # Fast path: read without locking (GIL-safe for dict.get)
        if not refresh:
            entry = _cache.get("summary")
            if entry and (time.time() - entry["ts"]) < _VCENTER_CACHE_TTL:
                return entry["data"]
        # Slow path: hold lock to prevent thundering herd on cache miss
        with _cache_lock:
            entry = _cache.get("summary")
            if not refresh and entry and (time.time() - entry["ts"]) < _VCENTER_CACHE_TTL:
                return entry["data"]
            data = _run(lambda c: c.get_summary())
            _cache["summary"] = {"data": data, "ts": time.time()}
            # Update vm→sap_number cache for customer name lookup
            try:
                from database import upsert_vm_sap_batch
                upsert_vm_sap_batch(data.get("vms", []), env_key)
            except Exception:
                pass
            return data

    @router.get("/hosts")
    def hosts(refresh: bool = False, request: Request = None):
        _check_admin_refresh(refresh, request)
        if not refresh:
            entry = _cache.get("summary")
            if entry and (time.time() - entry["ts"]) < _VCENTER_CACHE_TTL:
                return {"hosts": entry["data"].get("hosts", []), "clusters": entry["data"].get("clusters", [])}
        with _cache_lock:
            entry = _cache.get("summary")
            if not refresh and entry and (time.time() - entry["ts"]) < _VCENTER_CACHE_TTL:
                return {"hosts": entry["data"].get("hosts", []), "clusters": entry["data"].get("clusters", [])}
            data = _run(lambda c: c.get_summary())
            _cache["summary"] = {"data": data, "ts": time.time()}
            return {"hosts": data.get("hosts", []), "clusters": data.get("clusters", [])}

    @router.get("/cache-status")
    def cache_status():
        entry = _cache.get("summary")
        if not entry:
            return {"cached": False}
        age = int(time.time() - entry["ts"])
        return {
            "cached": True,
            "age_seconds": age,
            "ttl_seconds": _VCENTER_CACHE_TTL,
            "expires_in": max(0, _VCENTER_CACHE_TTL - age),
            "total_vms": entry["data"].get("total_vms"),
        }

    @router.get("/srm")
    def srm_info():
        return _run(lambda c: c.get_srm_info())

    @router.get("/vm-notes")
    def get_vm_notes_ep():
        return get_vm_notes(env_key)

    @router.put("/vm-notes")
    def update_vm_note_ep(entry: VmNoteIn):
        save_vm_note(env_key, entry.vm_name, entry.note)
        return {"status": "saved", "vm_name": entry.vm_name}

    @router.get("/vm-classification")
    def get_vm_classification_ep():
        return get_vm_classifications(env_key)

    @router.put("/vm-classification")
    def update_vm_classification_ep(entry: VmClassificationIn):
        save_vm_classification(env_key, entry.vm_name, entry.classification)
        return {"status": "saved", "vm_name": entry.vm_name}

    return router


app.include_router(build_vcenter_router("/api/vcenter/pt",     VCENTER_PT_HOST,     VCENTER_PT_USER,     VCENTER_PT_PASS))
app.include_router(build_vcenter_router("/api/vcenter/tc",     VCENTER_TC_HOST,     VCENTER_TC_USER,     VCENTER_TC_PASS))
app.include_router(build_vcenter_router("/api/vcenter/tc8",    VCENTER_TC8_HOST,    VCENTER_TC8_USER,    VCENTER_TC8_PASS))
app.include_router(build_vcenter_router("/api/vcenter/fattal", VCENTER_FATTAL_HOST, VCENTER_FATTAL_USER, VCENTER_FATTAL_PASS))


# ═══════════════════════════════════════════════════════════════════════════
# vCenter change-tracking: background polling + REST API
# ═══════════════════════════════════════════════════════════════════════════

_VC_ENVS_FOR_EVENTS = [
    ("TC",     lambda: (VCENTER_TC_HOST,  VCENTER_TC_USER,  VCENTER_TC_PASS)),
    ("TC8",    lambda: (VCENTER_TC8_HOST, VCENTER_TC8_USER, VCENTER_TC8_PASS)),
    ("PT",     lambda: (VCENTER_PT_HOST,  VCENTER_PT_USER,  VCENTER_PT_PASS)),
    ("Fattal", lambda: (VCENTER_FATTAL_HOST, VCENTER_FATTAL_USER, VCENTER_FATTAL_PASS)),
]

_events_poll_lock = threading.Lock()
_events_last_poll: dict = {}   # env -> epoch


def _do_poll_env(env_name: str, host: str, user: str, pwd: str, hours: int) -> int:
    if not host or not user:
        return 0
    try:
        events = VCenterClient(host=host, username=user, password=pwd).get_events(hours=hours)
        n = insert_vcenter_events(events, env_name)
        print(f"[vc-events] {env_name}: +{n} new events (polled {len(events)} raw, {hours}h)")
        return n
    except Exception as exc:
        print(f"[vc-events] {env_name} error: {exc}")
        return 0


def _poll_all_events(hours: int = 25):
    if not _events_poll_lock.acquire(blocking=False):
        return   # another poll running
    try:
        for env_name, cfg_fn in _VC_ENVS_FOR_EVENTS:
            h, u, p = cfg_fn()
            _do_poll_env(env_name, h, u, p, hours)
            _events_last_poll[env_name] = time.time()
    finally:
        _events_poll_lock.release()


def _polling_daemon():
    """Background thread: initial 30-day pull then every 2 h."""
    import time as _t
    _t.sleep(10)                          # let service fully start
    _poll_all_events(hours=720)           # initial: 30 days
    while True:
        _t.sleep(7200)                    # 2-hour cycle
        _poll_all_events(hours=25)


# Start daemon thread once — guard prevents duplicate threads on uvicorn --reload
if not any(t.name == "vc-events-poll" for t in threading.enumerate()):
    threading.Thread(target=_polling_daemon, daemon=True, name="vc-events-poll").start()


# ═══════════════════════════════════════════════════════════════════════════
# vCenter Performance page: per-minute CPU/RAM sampling (rolled up hourly/daily) + REST API
# ═══════════════════════════════════════════════════════════════════════════

import perf_store
import perf_collector

perf_store.init_perf_tables()
perf_collector.start(_VC_ENVS_FOR_EVENTS)


@app.on_event("shutdown")
def _perf_flush_on_shutdown():
    perf_collector.shutdown()


import atexit
atexit.register(perf_collector.shutdown)   # covers a stop that bypasses the ASGI shutdown event

_PERF_RANGES = {   # range key -> (days back, table). Hourly data is kept 30 days, daily ~6 months.
    "24h":  (1,   "vm_perf_hourly"),
    "7d":   (7,   "vm_perf_daily"),
    "30d":  (30,  "vm_perf_daily"),
    "90d":  (90,  "vm_perf_daily"),
    "180d": (183, "vm_perf_daily"),
}


def _perf_since(range_key: str):
    from datetime import datetime, timedelta
    days, table = _PERF_RANGES.get(range_key, _PERF_RANGES["7d"])
    if table == "vm_perf_hourly":
        since = (datetime.now() - timedelta(hours=24)).strftime("%Y-%m-%d %H:00")
    else:
        since = (datetime.now() - timedelta(days=days - 1)).strftime("%Y-%m-%d")
    return since, table


@app.get("/api/vcenter/perf")
def api_vcenter_perf(range: str = "7d", env: str | None = None):
    """Per-VM peak/avg CPU% and RAM% over the range, with when each peak happened."""
    since, table = _perf_since(range)
    return {
        "range": range,
        "since": since,
        "data":  perf_store.get_summary(since, table, env),
        "stats": perf_store.get_stats(),
    }


@app.get("/api/vcenter/perf/vm")
def api_vcenter_perf_vm(env: str, name: str, range: str = "7d"):
    """Time series for one VM: hourly buckets for up to 30 days, daily beyond that."""
    days, table = _PERF_RANGES.get(range, _PERF_RANGES["7d"])
    from datetime import datetime, timedelta
    if range in ("24h", "7d", "30d"):
        table = "vm_perf_hourly"
        since = (datetime.now() - timedelta(days=days)).strftime("%Y-%m-%d %H:00")
    else:
        since = (datetime.now() - timedelta(days=days - 1)).strftime("%Y-%m-%d")
    return {"table": table, "data": perf_store.get_series(env, name, since, table)}


# ── VM backup markers for the perf chart ─────────────────────────────────────
_bk_obj_cache: dict = {}   # server key -> (ts, {lower vm name -> object id})
_bk_pts_cache: dict = {}   # (server key, vm) -> (ts, [points])
_bk_map_cache: dict = {}   # server key -> (ts, {backupId -> backup entry})
_BK_OBJ_TTL = 900
_BK_PTS_TTL = 600
_BK_DEADLINE = 40   # seconds for the whole lookup (cold cache on a big server can be slow)


def _bk_clients():
    return {"PT": _pt_client, "Fattal": _fattal_client, "TC": _tc_client, "Lifesaver": _lifesaver_client}


def _bk_name_variants(name: str) -> list:
    """Lookup keys for a VM name, most specific first: the full name, the name without a trailing
    description in parentheses ('RKT-DB (Rakuto Priority)' -> 'rkt-db'), and the short host name
    without a DNS domain ('srv01.corp.local' -> 'srv01'). Veeam may know the VM or its agent under any of them."""
    full = " ".join((name or "").lower().split())
    out = [full]
    cut = full.find("(")
    base = full[:cut].strip() if cut > 0 else full
    for v in (base, full.split(".")[0].strip(), base.split(".")[0].strip()):
        if v and v not in out:
            out.append(v)
    return out


def _bk_object_ids(key: str, client) -> dict:
    import time as _t
    hit = _bk_obj_cache.get(key)
    if hit and _t.time() - hit[0] < _BK_OBJ_TTL:
        return hit[1]
    ids: dict = {}
    objs = [o for o in client.get_backup_objects(5000) if o.get("name") and o.get("id")]
    for o in objs:
        ids.setdefault(" ".join(o["name"].lower().split()), o["id"])
    for o in objs:                       # aliases: without a trailing "(...)" / without the DNS domain
        for v in _bk_name_variants(o["name"])[1:]:
            ids.setdefault(v, o["id"])
    _bk_obj_cache[key] = (_t.time(), ids)
    return ids


def _bk_backups_map(key: str, client) -> dict:
    import time as _t
    hit = _bk_map_cache.get(key)
    if hit and _t.time() - hit[0] < _BK_OBJ_TTL:
        return hit[1]
    m = client.get_backups_map()
    _bk_map_cache[key] = (_t.time(), m)
    return m


_REPLICA_RE = re.compile(r"snap|replica|(?<![a-z])rep(?![a-z])|target", re.I)


def _bk_local(iso: str):
    """ISO timestamp (any offset) -> naive local datetime, or None."""
    try:
        t = datetime.fromisoformat((iso or "").replace("Z", "+00:00"))
        return t.astimezone().replace(tzinfo=None) if t.tzinfo else t
    except Exception:
        return None


_bk_sess_cache: dict = {}   # server key -> (ts, [(name, start, end, id)])
_bk_task_cache: dict = {}   # (server key, session id) -> (ts, {vm lower -> (start, end)})


def _bk_sessions(key: str, client, since) -> list:
    import time as _t
    hit = _bk_sess_cache.get(key)
    if hit and _t.time() - hit[0] < _BK_PTS_TTL:
        return hit[1]
    raw = client.get_raw_sessions((since.astimezone() - timedelta(hours=6)).isoformat())
    out = []
    for s in raw:
        st, en = _bk_local(s.get("creationTime")), _bk_local(s.get("endTime"))
        if st:
            out.append(((s.get("name") or "").lower(), st, en or st, s.get("id", "")))
    _bk_sess_cache[key] = (_t.time(), out)
    return out


def _bk_window(key: str, client, vm: str, job: str, t, sessions: list):
    """(start, end) of the backup that produced the restore point created at t, or (None, None)."""
    import time as _t
    jl = job.lower()
    slack = timedelta(minutes=3)
    for name, st, en, sid in sessions:
        if not name or name not in jl:
            continue
        if not (st - slack <= t <= en + slack):
            continue
        ck = (key, sid)
        hit = _bk_task_cache.get(ck)
        if hit and _t.time() - hit[0] < _BK_PTS_TTL:
            tasks = hit[1]
        else:
            tasks = {}
            for ts in client.get_task_sessions(sid):
                a_, b_ = _bk_local(ts.get("creationTime")), _bk_local(ts.get("endTime"))
                if a_ and b_:
                    tasks[(ts.get("name") or "").lower()] = (a_, b_)
            _bk_task_cache[ck] = (_t.time(), tasks)
        w = tasks.get(vm.lower()) or tasks.get(vm.lower().split(".")[0])
        return w if w else (st, en)
    return (None, None)


def _bk_mark_frequent_as_replica(pts: list) -> None:
    """A job whose points come (median) <= 90 min apart, >= 6 of them, is a replication/snapshot job
    even when its name doesn't say so. Backups here run at most a few times a day."""
    by_job: dict = {}
    for p in pts:
        if p.get("kind") == "backup":
            t = _bk_local(p.get("time"))
            if t:
                by_job.setdefault(p.get("job", ""), []).append(t)
    for job, ts in by_job.items():
        if len(ts) < 6:
            continue
        ts.sort()
        gaps = sorted((b - a).total_seconds() for a, b in zip(ts, ts[1:]))
        if gaps[len(gaps) // 2] <= 90 * 60:
            for p in pts:
                if p.get("job", "") == job:
                    p["kind"] = "replica"


def _bk_kind(b: dict) -> str:
    """'replica' when the owning job looks like a replication/snapshot job (hourly VM replica points), else 'backup'.
    The Veeam REST restore point carries no type, so this goes by job name, e.g. 'FT-Rep-1', 'FT-Rep-1-target'."""
    return "replica" if _REPLICA_RE.search(str(b.get("name", ""))) else "backup"


@app.get("/api/vcenter/perf/backups")
def api_vcenter_perf_backups(name: str, days: int = 7, debug: bool = False):
    """Backup restore-point times of one VM across all Veeam servers (local time), for chart markers."""
    import time as _t
    from datetime import datetime, timedelta
    days = max(1, min(days, 183))
    since = datetime.now() - timedelta(days=days)
    variants = _bk_name_variants(name)
    out, errors, raw, matched = [], {}, {}, {}

    def _one(key, factory):
        o = []
        try:
            ck = (key, name.lower())
            hit = _bk_pts_cache.get(ck)
            if hit and _t.time() - hit[0] < _BK_PTS_TTL:
                pts = hit[1]
            else:
                cli = factory()
                ids = _bk_object_ids(key, cli)
                oid = None
                for v in variants:
                    if v in ids:
                        oid, matched[key] = ids[v], v
                        break
                pts = cli.get_restore_points(oid) if oid else []
                if pts:
                    bmap = _bk_backups_map(key, cli)
                    for p in pts:
                        b = bmap.get(p.get("backup_id")) or {}
                        p["job"] = b.get("name", "")
                        p["kind"] = _bk_kind(b)
                        p["backup_raw"] = b
                    _bk_mark_frequent_as_replica(pts)
                    sess = None
                    for p in pts:
                        if p.get("kind") != "backup":
                            continue
                        t = _bk_local(p.get("time"))
                        if not t:
                            continue
                        if sess is None:
                            sess = _bk_sessions(key, cli, since)
                        s_, e_ = _bk_window(key, cli, name, p.get("job", ""), t, sess)
                        if s_:
                            p["start"], p["end"] = s_.strftime("%Y-%m-%d %H:%M:%S"), e_.strftime("%Y-%m-%d %H:%M:%S")
                _bk_pts_cache[ck] = (_t.time(), pts)
            if debug and pts:
                raw[key] = {"point": pts[0].get("raw"), "backup": pts[0].get("backup_raw")}
            for p in pts:
                try:
                    t = datetime.fromisoformat(p["time"].replace("Z", "+00:00"))
                    t = t.astimezone().replace(tzinfo=None) if t.tzinfo else t
                except Exception:
                    continue
                if t >= since:
                    o.append({"server": key, "time": t.strftime("%Y-%m-%d %H:%M"), "type": p.get("type", ""),
                                "kind": p.get("kind", "backup"), "job": p.get("job", ""),
                                "start": p.get("start", ""), "end": p.get("end", "")})
        except Exception as e:
            errors[key] = str(e)[:120]
        return o

    # all servers in parallel, with an overall deadline so a dead server can't hang the chart
    import concurrent.futures as _cf
    ex = _cf.ThreadPoolExecutor(max_workers=4)
    futs = {ex.submit(_one, k, f): k for k, f in _bk_clients().items()}
    done, pending = _cf.wait(futs, timeout=_BK_DEADLINE)
    for f in done:
        out.extend(f.result())
    for f in pending:
        errors[futs[f]] = "timeout"
    ex.shutdown(wait=False, cancel_futures=True)
    out.sort(key=lambda x: x["time"])
    types: dict = {}
    for p in out:
        k = (p["server"] + " | " + p.get("job", "") + " | " + p["kind"]); types[k] = types.get(k, 0) + 1
    res = {"vm": name, "points": out, "types": types, "errors": errors}
    if debug:
        res["raw_sample"] = raw
        res["matched_by"] = matched
        res["tried"] = variants
    return res


@app.get("/api/vcenter/perf/hosts")
def api_vcenter_perf_hosts():
    """Live (refreshed every minute, not stored): which ESXi each VM runs on and how loaded that host is."""
    import time as _t
    now = _t.time()
    envs = {}
    for env, d in list(perf_collector.live.items()):
        envs[env] = {"age_s": int(now - d["ts"]), "hosts": d["hosts"], "vm_host": d["vm_host"]}
    return {"envs": envs}


_disk_cache: dict = {}   # (env, vm) -> (ts, result)
_DISK_LAT_WARN, _DISK_LAT_CRIT = 20, 30   # ms (VMware guidance: sustained >20 ms is slow, >30 ms is bad)


@app.get("/api/vcenter/perf/disks")
def api_vcenter_perf_disks(env: str, name: str):
    """Live per-virtual-disk latency / IOPS / throughput of one VM (read straight from vCenter, cached 30 s)."""
    import time as _t
    from vcenter_client import collect_vm_disks, open_session
    from pyVim.connect import Disconnect
    ck = (env, name.lower())
    hit = _disk_cache.get(ck)
    if hit and _t.time() - hit[0] < 30:
        return hit[1]
    cfg = next((fn for e_, fn in _VC_ENVS_FOR_EVENTS if e_ == env), None)
    if cfg is None:
        raise HTTPException(status_code=404, detail="Unknown environment")
    host, user, pwd = cfg()
    si = None
    try:
        si = open_session(host, user, pwd)
        res = collect_vm_disks(si, name)
    except Exception as e:
        raise HTTPException(status_code=502, detail="vCenter communication error")
    finally:
        if si is not None:
            try:
                Disconnect(si)
            except Exception:
                pass
    for d in res["disks"]:
        worst = max(d["read_ms_max"] or 0, d["write_ms_max"] or 0)
        avg = max(d["read_ms"] or 0, d["write_ms"] or 0)
        d["level"] = "crit" if avg >= _DISK_LAT_CRIT else ("warn" if avg >= _DISK_LAT_WARN or worst >= _DISK_LAT_CRIT else "ok")
    res["warn_ms"], res["crit_ms"] = _DISK_LAT_WARN, _DISK_LAT_CRIT
    _disk_cache[ck] = (_t.time(), res)
    return res


@app.get("/api/vcenter/perf/diagnose")
def api_vcenter_perf_diagnose(env: str, name: str, range: str = "7d"):
    """Rule-based diagnosis of one VM (metrics + live ESXi + live disk latency + backup timing)."""
    import concurrent.futures as _cf
    import perf_diagnose
    since, table = _perf_since(range)
    rows = perf_store.get_summary(since, table, env, name)
    summary = rows[0] if rows else None
    host = None
    live = perf_collector.live.get(env)
    if live:
        hn = live["vm_host"].get(name)
        host = live["hosts"].get(hn) if hn else None
    days = 1 if range == "24h" else int("".join(ch for ch in range if ch.isdigit()) or 7)
    disks = backups = None
    ex = _cf.ThreadPoolExecutor(max_workers=2)
    fd = ex.submit(api_vcenter_perf_disks, env, name)
    fb = ex.submit(api_vcenter_perf_backups, name, max(1, min(days, 30)), False)
    for fut, which in ((fd, "d"), (fb, "b")):
        try:
            res = fut.result(timeout=45)
            if which == "d":
                disks = res
            else:
                backups = res
        except Exception:
            pass
    ex.shutdown(wait=False, cancel_futures=True)
    return perf_diagnose.diagnose(env, name, range, summary, host, disks, backups)


class DiagFeedbackIn(BaseModel):
    env: str
    name: str
    range: str = ""
    rating: str                    # correct | partial | wrong
    area: str = ""
    comment: str = ""
    result: dict = {}


@app.post("/api/vcenter/perf/diagnose/feedback")
def api_vcenter_perf_diagnose_feedback(payload: DiagFeedbackIn, user: dict = Depends(get_current_user)):
    """Stores the user's verdict on a diagnosis (correct / partial / wrong + comment) to calibrate the rules."""
    import json as _json
    if payload.rating not in ("correct", "partial", "wrong"):
        raise HTTPException(status_code=400, detail="Invalid rating")
    blob = _json.dumps(payload.result, ensure_ascii=False)
    if len(blob) > 40000:
        raise HTTPException(status_code=413, detail="Result too large")
    fid = perf_store.save_feedback(str(user.get("username", "")), payload.env[:40], payload.name[:200], payload.range[:10],
                                   payload.rating, payload.area[:60], payload.comment[:1000], blob)
    return {"status": "ok", "id": fid}


@app.get("/api/vcenter/perf/diagnose/feedback", dependencies=[Depends(require_admin)])
def api_vcenter_perf_diagnose_feedback_list(format: str = "json", limit: int = 200):
    """Admin: the stored feedback; format=text is a compact digest to review and calibrate the rules from."""
    import json as _json
    from fastapi.responses import PlainTextResponse
    rows = perf_store.list_feedback(max(1, min(limit, 500)))
    if format != "text":
        return {"data": rows}
    counts, wrong_area = {}, {}
    for r in rows:
        counts[r["rating"]] = counts.get(r["rating"], 0) + 1
        if r["rating"] != "correct":
            wrong_area[r["area"] or "(general)"] = wrong_area.get(r["area"] or "(general)", 0) + 1
    lines = ["Diagnosis feedback: %s" % counts, "Not-correct by area: %s" % wrong_area, ""]
    for r in rows:
        try:
            res = _json.loads(r["result_json"] or "{}")
        except Exception:
            res = {}
        facts = res.get("facts", {})
        lines.append("#%s %s %s | %s/%s | range %s | rating=%s area=%s" % (r["id"], r["ts"], r["username"], r["environment"], r["vm_name"], r["range_key"], r["rating"], r["area"] or "-"))
        if r["comment"]:
            lines.append("  comment: %s" % r["comment"])
        lines.append("  verdict: %s" % ((res.get("verdict") or {}).get("en", "")))
        if res.get("cpu_answer"):
            lines.append("  cpu answer: %s" % res["cpu_answer"].get("en", ""))
        for f in res.get("findings", []):
            lines.append("  - [%s] %s: %s" % (f.get("sev"), f.get("area"), (f.get("title") or {}).get("en", "")))
        lines.append("  facts: %s" % _json.dumps(facts, ensure_ascii=False))
        lines.append("")
    return PlainTextResponse("\n".join(lines), media_type="text/plain; charset=utf-8")


@app.get("/api/vcenter/perf/rightsizing")
def api_vcenter_perf_rightsizing(env: str | None = None, days: int = 30, min_days: int = 14,
                                 cpu_target: float = 70, mem_target: float = 80):
    """Add/remove vCPU and RAM recommendations from the recurring (P95-of-daily) peaks."""
    import perf_rightsizing
    return perf_rightsizing.get_recommendations(
        env=env, days=max(7, min(days, 183)), min_days=max(1, min(min_days, 183)),
        cpu_target=max(30, min(cpu_target, 95)), mem_target=max(30, min(mem_target, 95)))


@app.get("/api/vcenter/perf/status")
def api_vcenter_perf_status():
    return {"collector": perf_collector.status, "sample_interval_s": perf_collector.SAMPLE_INTERVAL,
            "flush_interval_s": perf_collector.FLUSH_INTERVAL, **perf_store.get_stats()}


@app.get("/api/vcenter/events")
def api_vcenter_events(
    env:     str | None = None,
    from_dt: str | None = None,
    to_dt:   str | None = None,
    limit:   int        = 2000,
):
    return {
        "data":   get_vcenter_events(environment=env, from_dt=from_dt, to_dt=to_dt, limit=limit),
        "counts": count_vcenter_events_by_env(),
    }


@app.get("/api/vcenter/events/poll")
def api_vcenter_events_poll(hours: int = 25, _: str = Depends(require_admin)):
    """Manually trigger a poll across all vCenters. hours=25 (default) or 720 (30 days)."""
    h = max(1, min(hours, 720))
    threading.Thread(target=_poll_all_events, args=(h,), daemon=True).start()
    return {"status": "poll started", "hours": h}


@app.get("/api/vcenter/events/export")
def api_vcenter_events_export(
    env:     str | None = None,
    from_dt: str | None = None,
    to_dt:   str | None = None,
):
    from fastapi.responses import StreamingResponse
    rows = get_vcenter_events(environment=env, from_dt=from_dt, to_dt=to_dt, limit=50000)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "vCenter Changes"
    ws.append(["\u05E1\u05D1\u05D9\u05D1\u05D4", "VM",
               "\u05D0\u05D9\u05E8\u05D5\u05E2",
               "\u05E4\u05E8\u05D8\u05D9\u05DD",
               "\u05DE\u05E9\u05EA\u05DE\u05E9",
               "\u05EA\u05D0\u05E8\u05D9\u05DA \u05D5\u05E9\u05E2\u05D4"])
    def _safe(v) -> str:
        """Prevent formula injection in Excel cells."""
        s = str(v) if v is not None else ""
        return ("'" + s) if s and s[0] in ('=', '+', '-', '@', '\t', '\r') else s

    for r in rows:
        detail = r.get("change_info") or r.get("description") or ""
        ws.append([
            _safe(r.get("environment", "")), _safe(r.get("vm_name", "")),
            _safe(r.get("event_label") or r.get("event_type", "")),
            _safe(detail), _safe(r.get("user_name", "")), _safe(r.get("timestamp", "")),
        ])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=vcenter_changes.xlsx"},
    )


# ═══════════════════════════════════════════════════════════════════════════
# vCenter History: on-demand full VM inventory snapshots
# ═══════════════════════════════════════════════════════════════════════════

VM_SNAPSHOT_RETENTION = int(os.getenv("VM_SNAPSHOT_RETENTION", "100"))
# Password required to delete a *protected* snapshot (e.g. the automatic monthly archive).
# Protected snapshots are excluded from VM_SNAPSHOT_RETENTION pruning entirely — see
# prune_old_snapshots() in database.py — so this is the only way they're ever removed.
SNAPSHOT_DELETE_PASSWORD = os.getenv("SNAPSHOT_DELETE_PASSWORD", "H@poelTLV")


def _take_vm_snapshot(taken_by: str, note: str = "", protected: bool = False) -> dict:
    snap_id = create_vm_snapshot(taken_by, note, protected=protected)
    ok_envs: list[str] = []
    total = 0

    def _fetch(env_name: str, host: str, user: str, pwd: str):
        if not host:
            return env_name, None
        data = VCenterClient(host=host, username=user, password=pwd).get_summary()
        return env_name, data

    with ThreadPoolExecutor(max_workers=len(_VC_ENVS_FOR_EVENTS)) as pool:
        futures = []
        for env_name, cfg_fn in _VC_ENVS_FOR_EVENTS:
            h, u, p = cfg_fn()
            futures.append(pool.submit(_fetch, env_name, h, u, p))
        for fut in as_completed(futures):
            try:
                env_name, data = fut.result()
            except Exception as exc:
                print(f"[vm-snapshot] error: {exc}")
                continue
            if data is None:
                continue
            vms = data.get("vms", [])
            try:
                insert_snapshot_items(snap_id, env_name, vms)
                # Full payload (vms + hosts + clusters + datastores + affinity_rules + counts) —
                # a superset of insert_snapshot_items(), for read-only consumers that need the
                # complete point-in-time picture (e.g. Hosts/Clusters, not just VM deltas).
                insert_snapshot_env_data(snap_id, env_name, data)
                total += len(vms)
                ok_envs.append(env_name)
            except Exception as exc:
                print(f"[vm-snapshot] {env_name} insert error: {exc}")

    finalize_snapshot(snap_id, total, ok_envs)
    pruned = prune_old_snapshots(keep=VM_SNAPSHOT_RETENTION)
    return {"id": snap_id, "vm_count": total, "envs_ok": ok_envs, "pruned": pruned}


@app.post("/api/vcenter/history/snapshot")
def api_take_vm_snapshot(note: str = "", user: str = Depends(require_admin)):
    """Synchronously pull all vCenters' current VM inventory and store it as a new version."""
    result = _take_vm_snapshot(taken_by=user, note=note)
    return result


@app.get("/api/vcenter/history/snapshots")
def api_list_vm_snapshots(limit: int = 100):
    return {"data": list_snapshots(limit=limit)}


@app.get("/api/vcenter/history/snapshot/{snapshot_id}")
def api_get_vm_snapshot(snapshot_id: int, environment: str | None = None):
    snap = get_snapshot(snapshot_id)
    if not snap:
        raise HTTPException(status_code=404, detail="Snapshot not found")
    return {"snapshot": snap, "vms": get_snapshot_items(snapshot_id, environment)}


class SnapshotDeleteIn(BaseModel):
    password: str = ""


@app.delete("/api/vcenter/history/snapshot/{snapshot_id}")
def api_delete_vm_snapshot(snapshot_id: int, payload: SnapshotDeleteIn = SnapshotDeleteIn(), _: str = Depends(require_admin)):
    snap = get_snapshot(snapshot_id)
    if not snap:
        raise HTTPException(status_code=404, detail="Snapshot not found")
    if snap.get("protected"):
        if not payload.password or payload.password != SNAPSHOT_DELETE_PASSWORD:
            raise HTTPException(status_code=403, detail="Incorrect password for a protected snapshot")
    if not delete_snapshot(snapshot_id):
        raise HTTPException(status_code=404, detail="Snapshot not found")
    return {"status": "deleted", "id": snapshot_id}


def _seconds_until_next_month_start() -> float:
    """Seconds from now until the next 1st-of-month at 00:01 local time. If it's the 1st and
    still before 00:01, targets *today*; otherwise rolls over to next month."""
    now = datetime.now()
    target = datetime(now.year, now.month, 1, 0, 1, 0)
    if target <= now:
        year, month = now.year, now.month + 1
        if month > 12:
            month = 1
            year += 1
        target = datetime(year, month, 1, 0, 1, 0)
    return max((target - now).total_seconds(), 1)


def _monthly_snapshot_daemon():
    """Background thread: takes a *protected* full VM-inventory snapshot on the 1st of every
    month at 00:01. Protected snapshots are excluded from VM_SNAPSHOT_RETENTION pruning
    (prune_old_snapshots() in database.py) and can only be deleted with SNAPSHOT_DELETE_PASSWORD
    (api_delete_vm_snapshot above) — they form a permanent monthly archive, separate from the
    regular rolling history."""
    while True:
        sleep_s = _seconds_until_next_month_start()
        print(f"[monthly-snapshot] sleeping {sleep_s/3600:.1f}h until next run")
        time.sleep(sleep_s)
        try:
            if has_protected_snapshot_today():
                print("[monthly-snapshot] a protected snapshot already exists for today — skipping")
                continue
            result = _take_vm_snapshot(taken_by="system-scheduled", note="גיבוי אוטומטי חודשי", protected=True)
            print(f"[monthly-snapshot] created snapshot id={result['id']} vm_count={result['vm_count']} envs_ok={result['envs_ok']}")
        except Exception as exc:
            print(f"[monthly-snapshot] error: {exc}")


# Guard prevents duplicate threads on uvicorn --reload, same pattern as vc-events-poll above.
if not any(t.name == "monthly-snapshot" for t in threading.enumerate()):
    threading.Thread(target=_monthly_snapshot_daemon, daemon=True, name="monthly-snapshot").start()


def _seconds_until_next_daily_run(hour: int = 1, minute: int = 15) -> float:
    """Seconds from now until the next daily run at HH:MM local time (default 01:15 — offset
    from the monthly job's 00:01 so the two don't fetch from all vCenters at the same instant).
    If that time already passed today, targets tomorrow."""
    now = datetime.now()
    target = datetime(now.year, now.month, now.day, hour, minute, 0)
    if target <= now:
        target = target + timedelta(days=1)
    return max((target - now).total_seconds(), 1)


def _daily_snapshot_daemon():
    """Background thread: takes a regular (non-protected) full VM-inventory snapshot once a day
    at 01:15. Counts against VM_SNAPSHOT_RETENTION like any manual snapshot — this is meant for
    external read-only consumers (e.g. another dashboard reading dashboard.db directly) that
    want a daily-fresh picture without needing live API access. Separate from the monthly
    *protected* archive above, which exists for a different purpose (permanent record)."""
    while True:
        sleep_s = _seconds_until_next_daily_run()
        print(f"[daily-snapshot] sleeping {sleep_s/3600:.1f}h until next run")
        time.sleep(sleep_s)
        try:
            if has_snapshot_today_by("system-daily"):
                print("[daily-snapshot] a daily snapshot already exists for today — skipping")
                continue
            result = _take_vm_snapshot(taken_by="system-daily", note="גיבוי יומי אוטומטי", protected=False)
            print(f"[daily-snapshot] created snapshot id={result['id']} vm_count={result['vm_count']} envs_ok={result['envs_ok']}")
        except Exception as exc:
            print(f"[daily-snapshot] error: {exc}")


if not any(t.name == "daily-snapshot" for t in threading.enumerate()):
    threading.Thread(target=_daily_snapshot_daemon, daemon=True, name="daily-snapshot").start()


# ═══════════════════════════════════════════════════════════════════════════
# Customer billing config: optional per-customer free-growth allowance + per-resource
# pricing, used by the History change report to compute billable resource growth.
# ═══════════════════════════════════════════════════════════════════════════

class BillingConfigIn(BaseModel):
    customer_name:    str
    enabled:           bool  = True
    free_pct_cpu:      float = 5
    free_pct_ram:      float = 5
    free_pct_stor:     float = 5
    price_vcpu:        float = 0
    price_ram_gb:      float = 0
    price_storage_gb:  float = 0


@app.get("/api/billing/config")
def api_list_billing_config(_: str = Depends(require_admin)):
    return {"data": list_billing_configs()}


@app.put("/api/billing/config")
def api_upsert_billing_config(payload: BillingConfigIn, _: str = Depends(require_admin)):
    name = payload.customer_name.strip()
    if not name:
        raise HTTPException(status_code=400, detail="customer_name is required")
    for field, val in (("free_pct_cpu", payload.free_pct_cpu), ("free_pct_ram", payload.free_pct_ram),
                        ("free_pct_stor", payload.free_pct_stor)):
        if val < 0 or val > 100:
            raise HTTPException(status_code=400, detail=f"{field} must be between 0 and 100")
    for field, val in (("price_vcpu", payload.price_vcpu), ("price_ram_gb", payload.price_ram_gb),
                        ("price_storage_gb", payload.price_storage_gb)):
        if val < 0:
            raise HTTPException(status_code=400, detail=f"{field} must not be negative")
    upsert_billing_config(
        name, payload.enabled, payload.free_pct_cpu, payload.free_pct_ram, payload.free_pct_stor,
        payload.price_vcpu, payload.price_ram_gb, payload.price_storage_gb,
    )
    return get_billing_config(name)


@app.delete("/api/billing/config/{customer_name}")
def api_delete_billing_config(customer_name: str, _: str = Depends(require_admin)):
    if not delete_billing_config(customer_name):
        raise HTTPException(status_code=404, detail="No billing config for this customer")
    return {"status": "deleted", "customer_name": customer_name}


# ------------------------------------------------------------------ #
#  Veeam Enterprise Manager                                            #
# ------------------------------------------------------------------ #

VEEAM_EM_HOST = os.getenv("VEEAM_EM_HOST", "al-veeambem.allegronet.co.il:9398")
VEEAM_PT_HOST = os.getenv("VEEAM_PT_HOST", "al-veeampt.allegronet.co.il:9419")
VEEAM_PT_USER = os.getenv("VEEAM_USER", r"al-lifesaver\VeeamRO")
VEEAM_FATTAL_HOST = os.getenv("VEEAM_FATTAL_HOST", "ft-vbck.fattal.co.il:9419")
VEEAM_FATTAL_USER = os.getenv("VEEAM_FATTAL_USER", r"al-lifesaver\VeeamRO")
VEEAM_LIFESAVER_HOST = os.getenv("VEEAM_LIFESAVER_HOST", "al-lifesaver.allegronet.co.il:9419")
VEEAM_LIFESAVER_USER = os.getenv("VEEAM_LIFESAVER_USER", r"al-lifesaver\VeeamRO")
VEEAM_TC_HOST = os.getenv("VEEAM_TC_HOST", "al-veeamcc.allegronet.co.il:9419")
VEEAM_TC_USER = os.getenv("VEEAM_TC_USER", r"al-veeamcc\VeeamRO")

VEEAM_EM_USER = os.getenv("VEEAM_USER",    r"al-lifesaver\VeeamRO")

_veeam_cache: dict = {}
_VEEAM_CACHE_TTL    = 120  # Veeam EM
_VEEAM_BR_CACHE_TTL = 300  # Veeam B&R (pt / fattal / tc / lifesaver)

def _veeam_client():
    return VeeamEMClient(host=VEEAM_EM_HOST, username=VEEAM_EM_USER)

veeam_router = APIRouter(prefix="/api/backup")

@veeam_router.get("/ping")
def veeam_ping():
    try:
        return _veeam_client().ping()
    except Exception as e:
        raise HTTPException(status_code=502, detail="Veeam connection error")

@veeam_router.get("/probe")
def veeam_probe():
    try:
        return _veeam_client().probe()
    except Exception as e:
        raise HTTPException(status_code=502, detail=str(e))

@veeam_router.get("/raw-replicas")
def veeam_raw_replicas():
    """Return raw replica + session data for field mapping diagnostics."""
    try:
        import requests as _req, urllib3 as _u3
        _u3.disable_warnings()
        import os as _os
        _CA = _os.getenv("INFRA_CA_BUNDLE", "").strip()
        _TLS = (True if _CA.lower() in ("true","1") else (_CA if _CA and _CA.lower() not in ("false","0") else False))
        c = _veeam_client()
        c._create_session()
        # Get first 3 unique replicas
        url = f"{c._base()}/query?type=Replica&format=Entities&pageSize=100"
        r = _req.get(url, headers=c._hdrs(), verify=_TLS, timeout=90)
        raw = r.json()
        reps = (raw.get("entities") or raw.get("Entities") or {}).get("replicas", {}).get("replicas", [])
        seen, unique = set(), []
        for rep in reps:
            uid = rep.get("uid","")
            if uid not in seen:
                seen.add(uid)
                unique.append(rep)
        # Also get one entity detail
        entity_detail = None
        if unique and unique[0].get("href"):
            try:
                er = _req.get(unique[0]["href"], headers=c._hdrs(), verify=_TLS, timeout=30)
                entity_detail = er.json()
            except Exception as ex:
                entity_detail = {"error": str(ex)}
        # Get sessions
        surl = f"{c._base()}/query?type=ReplicaJobSession&format=Entities&pageSize=5"
        sr = _req.get(surl, headers=c._hdrs(), verify=_TLS, timeout=90)
        sessions_raw = sr.json()
        return {
            "unique_replicas_sample": unique[:3],
            "entity_detail_sample": entity_detail,
            "sessions_raw": sessions_raw,
        }
    except Exception as e:
        raise HTTPException(status_code=502, detail=str(e))

@veeam_router.get("/summary")
def veeam_summary(refresh: bool = False, request: Request = None):
    _check_admin_refresh(refresh, request)
    entry = _veeam_cache.get("summary")
    if entry and not refresh and (time.time() - entry["ts"]) < _VEEAM_CACHE_TTL:
        return entry["data"]
    try:
        data = _veeam_client().get_summary()
        _veeam_cache["summary"] = {"data": data, "ts": time.time()}
        return data
    except Exception as e:
        raise HTTPException(status_code=502, detail="Veeam connection error")

@veeam_router.get("/dr")
def veeam_dr(request: Request = None):
    entry = _veeam_cache.get("dr")
    if entry and (time.time() - entry["ts"]) < _VEEAM_CACHE_TTL:
        return entry["data"]
    try:
        data = _veeam_client().get_dr_summary()
        _veeam_cache["dr"] = {"data": data, "ts": time.time()}
        return data
    except Exception as e:
        raise HTTPException(status_code=502, detail="Veeam DR connection error")

app.include_router(veeam_router)

def _br_probe(client_fn):
    """Shared B&R probe: check jobs (with type), sessions, repos."""
    import requests as _req, urllib3 as _u3
    _u3.disable_warnings()
    import os as _os
    _CA = _os.getenv("INFRA_CA_BUNDLE","").strip()
    _TLS = (True if _CA.lower() in ("true","1") else (_CA if _CA and _CA.lower() not in ("false","0") else False))
    try:
        c = client_fn()
        c._get_token()
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"Auth failed: {e}")
    results = {"host": c.host, "auth": "ok", "endpoints": {}}
    for path in ["jobs?limit=200", "sessions?limit=50",
                 "backupInfrastructure/repositories?limit=10"]:
        url = f"{c.host}/api/v1/{path}"
        try:
            r = _req.get(url, headers=c._hdrs(), verify=_TLS, timeout=45)
            try:
                body = r.json()
                data = body.get("data", body) if isinstance(body, dict) else body
                total = body.get("pagination", {}).get("total", len(data)) if isinstance(body, dict) else len(data)
                results["endpoints"][path] = {
                    "status": r.status_code,
                    "total": total,
                    "sample": data[:5] if isinstance(data, list) else str(data)[:300],
                }
            except Exception:
                results["endpoints"][path] = {"status": r.status_code, "raw": r.text[:300]}
        except Exception as ex:
            results["endpoints"][path] = {"error": str(ex)}
    return results

# ── PT Backup (Veeam B&R, port 9419) ──────────────────────────────────────── #
_pt_cache: dict = {}

def _pt_client():
    return VeeamBRClient(host=VEEAM_PT_HOST, username=VEEAM_PT_USER)

pt_backup_router = APIRouter(prefix="/api/backup/pt")

@pt_backup_router.get("/ping")
def pt_backup_ping():
    try:
        return _pt_client().ping()
    except Exception as e:
        raise HTTPException(status_code=502, detail="Backup server communication error")

@pt_backup_router.get("/probe")
def pt_backup_probe():
    return _br_probe(_pt_client)

@pt_backup_router.get("/summary")
def pt_backup_summary(refresh: bool = False, request: Request = None):
    _check_admin_refresh(refresh, request)
    entry = _pt_cache.get("summary")
    if entry and not refresh and (time.time() - entry["ts"]) < _VEEAM_BR_CACHE_TTL:
        return entry["data"]
    try:
        data = _pt_client().get_summary()
        _pt_cache["summary"] = {"data": data, "ts": time.time()}
        return data
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"PT backup communication error: {e}")

@pt_backup_router.get("/objects")
def pt_backup_objects():
    try:
        return {"data": _pt_client().get_backup_objects()}
    except Exception as e:
        raise HTTPException(status_code=502, detail="Backup server communication error")

app.include_router(pt_backup_router)

_fattal_cache: dict = {}

def _fattal_client():
    pw = os.getenv("VEEAM_FATTAL_PASS") or os.getenv("VEEAM_PASS", "")
    return VeeamBRClient(host=VEEAM_FATTAL_HOST, username=VEEAM_FATTAL_USER, password=pw)

fattal_backup_router = APIRouter(prefix="/api/backup/fattal")

@fattal_backup_router.get("/ping")
def fattal_backup_ping():
    try:
        return _fattal_client().ping()
    except Exception as e:
        raise HTTPException(status_code=502, detail="Backup server communication error")

@fattal_backup_router.get("/probe")
def fattal_backup_probe():
    return _br_probe(_fattal_client)

@fattal_backup_router.get("/summary")
def fattal_backup_summary(refresh: bool = False, request: Request = None):
    _check_admin_refresh(refresh, request)
    entry = _fattal_cache.get("summary")
    if entry and not refresh and (time.time() - entry["ts"]) < _VEEAM_BR_CACHE_TTL:
        return entry["data"]
    try:
        data = _fattal_client().get_summary()
        _fattal_cache["summary"] = {"data": data, "ts": time.time()}
        return data
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"Fattal backup communication error: {e}")

@fattal_backup_router.get("/objects")
def fattal_backup_objects():
    try:
        return {"data": _fattal_client().get_backup_objects()}
    except Exception as e:
        raise HTTPException(status_code=502, detail="Backup server communication error")

app.include_router(fattal_backup_router)
_tc_cache: dict = {}

def _tc_client():
    pw = os.getenv("VEEAM_TC_PASS", "")
    return VeeamBRClient(host=VEEAM_TC_HOST, username=VEEAM_TC_USER, password=pw)

tc_backup_router = APIRouter(prefix="/api/backup/tc")

@tc_backup_router.get("/ping")
def tc_backup_ping():
    try:
        return _tc_client().ping()
    except Exception as e:
        raise HTTPException(status_code=502, detail="Backup server communication error")

@tc_backup_router.get("/probe")
def tc_backup_probe():
    """Raw probe of TC Veeam B&R API endpoints including Cloud Connect."""
    try:
        import requests as _req, urllib3 as _u3
        _u3.disable_warnings()
        import os as _os
        _CA = _os.getenv("INFRA_CA_BUNDLE","").strip()
        _TLS = (True if _CA.lower() in ("true","1") else (_CA if _CA and _CA.lower() not in ("false","0") else False))
        c = _tc_client()
        c._get_token()
        results = {"host": c.host, "auth": "ok"}
        paths = [
            "jobs?limit=5",
            "sessions?limit=5",
            "cloudConnect/replicas?limit=20",
            "cloudConnect/tenants?limit=20",
            "cloudConnect?limit=5",
            "replicas?limit=20",
            "backupObjects?limit=5",
        ]
        for path in paths:
            url = f"{c.host}/api/v1/{path}"
            try:
                r = _req.get(url, headers=c._hdrs(), verify=_TLS, timeout=30)
                try: body = r.json()
                except: body = r.text[:300]
                data = body.get("data", body) if isinstance(body, dict) else body
                total = body.get("pagination", {}).get("total", "") if isinstance(body, dict) else ""
                results[path] = {"status": r.status_code, "total": total, "sample": str(data)[:500]}
            except Exception as ex:
                results[path] = {"error": str(ex)}
        return results
    except Exception as e:
        raise HTTPException(status_code=502, detail=str(e))

@tc_backup_router.get("/cloud-replicas")
def tc_cloud_replicas():
    """Fetch Cloud Connect tenant replica VMs (external clients replicating to us)."""
    try:
        import requests as _req, urllib3 as _u3
        _u3.disable_warnings()
        import os as _os
        _CA = _os.getenv("INFRA_CA_BUNDLE","").strip()
        _TLS = (True if _CA.lower() in ("true","1") else (_CA if _CA and _CA.lower() not in ("false","0") else False))
        c = _tc_client()
        c._get_token()

        # ── Detection rules ──────────────────────────────────────────────────
        # Prefix rules: name starts with key (case-insensitive) → (client, prefix)
        PREFIX_RULES = {"tadiran_": "Tadiran"}

        # Cluster rules: cluster == key AND name ends with _replica → client
        CLUSTER_RULES = {"INF-TC-ISROTEL_ibox03": "Isrotel"}

        # Offline fallback: Offline_VMs cluster, name prefix → client
        OFFLINE_RULES = {
            "ELT":    "Isrotel",
            "IS-OPR": "Isrotel",
            "IS-MDT": "Isrotel",
            "IS-OPT": "Isrotel",
        }

        SUFFIX_STRIP = ("_replica", "_copy", "_restored")

        def _detect(nm, cl):
            nl = nm.lower()
            for pfx, cn in PREFIX_RULES.items():
                if nl.startswith(pfx):
                    return cn, pfx
            if cl in CLUSTER_RULES and nl.endswith("_replica"):
                return CLUSTER_RULES[cl], ""
            if cl == "Offline_VMs" and nl.endswith("_replica"):
                for pfx, cn in OFFLINE_RULES.items():
                    if nm.upper().startswith(pfx.upper()):
                        return cn, ""
            return None, ""

        def _clean(nm, pfx_strip):
            n = nm[len(pfx_strip):] if pfx_strip and nm.lower().startswith(pfx_strip) else nm
            for sfx in SUFFIX_STRIP:
                if n.lower().endswith(sfx):
                    n = n[:-len(sfx)]
                    break
            return n.strip(" -_") or nm

        # Paginate all backupObjects
        all_objects = []
        skip = 0
        while True:
            r = _req.get(f"{c.host}/api/v1/backupObjects?limit=500&skip={skip}",
                         headers=c._hdrs(), verify=_TLS, timeout=60)
            if r.status_code != 200:
                break
            body = r.json()
            page = body.get("data", []) if isinstance(body, dict) else []
            all_objects.extend(page)
            total = body.get("pagination", {}).get("total", 0) if isinstance(body, dict) else 0
            skip += 500
            if skip >= total or not page:
                break

        # Filter to client VMs, deduplicate by (client, clean_name)
        clients: dict = {}
        seen: set = set()
        for obj in all_objects:
            name = obj.get("name", "")
            if obj.get("type") != "VM" or not name:
                continue
            path = obj.get("path", "")
            parts = path.split("/")
            cluster = parts[1] if len(parts) > 1 else ""

            client_name, pfx_strip = _detect(name, cluster)
            if not client_name:
                continue
            clean = _clean(name, pfx_strip)
            key = (client_name, clean)
            if key in seen:
                continue
            seen.add(key)

            clients.setdefault(client_name, []).append({
                "id":                 obj.get("id", ""),
                "name":               name,
                "clean_name":         clean,
                "path":               path,
                "cluster":            cluster,
                "restore_points":     0,
                "last_restore_point": "",
                "status":             "unknown",
            })

        # Fetch latest restore point for each VM (gives us creation date + count)
        for vms in clients.values():
            for vm in vms:
                if not vm["id"]:
                    continue
                try:
                    rp_r = _req.get(
                        f"{c.host}/api/v1/backupObjects/{vm['id']}/restorePoints"
                        "?limit=1&orderColumn=CreationTime&orderAsc=false",
                        headers=c._hdrs(), verify=_TLS, timeout=20
                    )
                    if rp_r.status_code == 200:
                        rb = rp_r.json()
                        rp_list = rb.get("data", []) if isinstance(rb, dict) else []
                        total_rp = rb.get("pagination", {}).get("total", 0) if isinstance(rb, dict) else 0
                        vm["restore_points"] = total_rp
                        if rp_list:
                            vm["last_restore_point"] = rp_list[0].get("creationTime", "")
                            vm["status"] = "ready"
                        else:
                            vm["status"] = "no_points"
                except Exception:
                    pass

        return {
            "clients":    clients,
            "total_vms":  sum(len(v) for v in clients.values()),
            "server":     c.host,
        }
    except Exception as e:
        raise HTTPException(status_code=502, detail=str(e))

@tc_backup_router.get("/scan-replicas")
def tc_scan_replicas():
    """Scan al-veeamcc sessions by type + check Lifesaver/PT backupObjects for replicas."""
    try:
        import requests as _req, urllib3 as _u3, os as _os
        _u3.disable_warnings()
        _CA = _os.getenv("INFRA_CA_BUNDLE","").strip()
        _TLS = (True if _CA.lower() in ("true","1") else (_CA if _CA and _CA.lower() not in ("false","0") else False))

        result = {}

        # 1. al-veeamcc: all sessions → breakdown by sessionType
        try:
            c_tc = _tc_client()
            c_tc._get_token()
            sess_types: dict = {}
            sess_skip = 0
            sess_total = 9999
            while sess_skip < sess_total:
                r = _req.get(f"{c_tc.host}/api/v1/sessions?limit=500&skip={sess_skip}",
                             headers=c_tc._hdrs(), verify=_TLS, timeout=30)
                if r.status_code != 200:
                    break
                body = r.json()
                for s in body.get("data", []):
                    t = s.get("sessionType", "unknown")
                    sess_types[t] = sess_types.get(t, 0) + 1
                sess_total = body.get("pagination", {}).get("total", 0)
                sess_skip += 500
                if not body.get("data"):
                    break
            result["cc_session_types"] = sess_types
        except Exception as e:
            result["cc_session_types"] = {"error": str(e)}

        # 2. Lifesaver + PT: backupObjects → find _replica VMs
        from veeam_br_client import VeeamBRClient
        _VEEAM_USER = _os.getenv("VEEAM_USER","")
        _VEEAM_PASS = _os.getenv("VEEAM_PASS","")
        _LS_HOST    = _os.getenv("LIFESAVER_HOST","al-lifesaver.allegronet.co.il:9419")
        _PT_HOST    = _os.getenv("VEEAM_PT_HOST", "al-veeampt.allegronet.co.il:9419")

        for label, host in [("lifesaver", _LS_HOST), ("pt", _PT_HOST)]:
            try:
                cli = VeeamBRClient(host, _VEEAM_USER, _VEEAM_PASS)
                cli._get_token()
                rep_vms = []
                skip2 = 0
                total2 = 9999
                while skip2 < total2:
                    r2 = _req.get(f"{cli.host}/api/v1/backupObjects?limit=500&skip={skip2}",
                                  headers=cli._hdrs(), verify=_TLS, timeout=30)
                    if r2.status_code != 200:
                        break
                    b2 = r2.json()
                    for obj in b2.get("data", []):
                        n = obj.get("name","")
                        if obj.get("type") == "VM" and n.lower().endswith("_replica"):
                            p = obj.get("path","")
                            rep_vms.append({"name": n, "path": p,
                                            "id": obj.get("id","")})
                    total2 = b2.get("pagination",{}).get("total",0)
                    skip2 += 500
                    if not b2.get("data"):
                        break
                result[f"{label}_replica_vms"] = rep_vms[:200]
                result[f"{label}_total_objects"] = total2
            except Exception as ex:
                result[f"{label}_replica_vms"] = []
                result[f"{label}_error"] = str(ex)

        return result
    except Exception as e:
        raise HTTPException(status_code=502, detail=str(e))


@tc_backup_router.get("/summary")
def tc_backup_summary(refresh: bool = False, request: Request = None):
    _check_admin_refresh(refresh, request)
    entry = _tc_cache.get("summary")
    if entry and not refresh and (time.time() - entry["ts"]) < _VEEAM_BR_CACHE_TTL:
        return entry["data"]
    try:
        data = _tc_client().get_summary()
        _tc_cache["summary"] = {"data": data, "ts": time.time()}
        return data
    except Exception as e:
        raise HTTPException(status_code=502, detail="Backup server communication error")

@tc_backup_router.get("/objects")
def tc_backup_objects():
    try:
        return {"data": _tc_client().get_backup_objects()}
    except Exception as e:
        raise HTTPException(status_code=502, detail="Backup server communication error")

app.include_router(tc_backup_router)
_lifesaver_cache: dict = {}

def _lifesaver_client():
    pw = os.getenv("VEEAM_LIFESAVER_PASS", "")
    return VeeamBRClient(host=VEEAM_LIFESAVER_HOST, username=VEEAM_LIFESAVER_USER, password=pw)

lifesaver_backup_router = APIRouter(prefix="/api/backup/lifesaver")

@lifesaver_backup_router.get("/ping")
def lifesaver_backup_ping():
    try:
        return _lifesaver_client().ping()
    except Exception as e:
        raise HTTPException(status_code=502, detail="Backup server communication error")

@lifesaver_backup_router.get("/probe")
def lifesaver_backup_probe():
    return _br_probe(_lifesaver_client)

@lifesaver_backup_router.get("/summary")
def lifesaver_backup_summary(refresh: bool = False, request: Request = None):
    _check_admin_refresh(refresh, request)
    entry = _lifesaver_cache.get("summary")
    if entry and not refresh and (time.time() - entry["ts"]) < _VEEAM_BR_CACHE_TTL:
        return entry["data"]
    try:
        data = _lifesaver_client().get_summary()
        _lifesaver_cache["summary"] = {"data": data, "ts": time.time()}
        return data
    except Exception as e:
        raise HTTPException(status_code=502, detail="Backup server communication error")

@lifesaver_backup_router.get("/objects")
def lifesaver_backup_objects():
    try:
        return {"data": _lifesaver_client().get_backup_objects()}
    except Exception as e:
        raise HTTPException(status_code=502, detail="Backup server communication error")

app.include_router(lifesaver_backup_router)





# ------------------------------------------------------------------ #
#  DR aggregated summary endpoint                                      #
# ------------------------------------------------------------------ #

_dr_cache: dict = {}
_DR_CACHE_TTL = 120  # seconds

@app.get("/api/dr/summary")
def dr_summary(refresh: bool = False, request: Request = None):
    """
    Aggregate DR data from all vCenters and Infinidat iBoxes in parallel.
    Pass ?refresh=true to bypass the cache — admin only.
    """
    _check_admin_refresh(refresh, request)
    _dr_entry = _dr_cache.get("summary")
    if not refresh and _dr_entry and (time.time() - _dr_entry["ts"]) < _DR_CACHE_TTL:
        return _dr_entry["data"]

    cust_map = {}
    try:
        cust_map = get_customers_map()
    except Exception:
        pass

    # Define all fetch tasks
    vc_configs = {
        "tc":     (VCENTER_TC_HOST,     VCENTER_TC_USER,     VCENTER_TC_PASS),
        "pt":     (VCENTER_PT_HOST,     VCENTER_PT_USER,     VCENTER_PT_PASS),
        "fattal": (VCENTER_FATTAL_HOST, VCENTER_FATTAL_USER, VCENTER_FATTAL_PASS),
    }
    ib_configs = {
        "tc03":   (IBOX_TC03_HOST,  IBOX_TC03_USER,  IBOX_TC03_PASS,  "TC iBox-03"),
        "tc_old": (IBOX_TC_OLD_HOST, IBOX_TC_OLD_USER, IBOX_TC_OLD_PASS, "TC iBox (Old)"),
        "pt":     (IBOX_PT_HOST,    IBOX_PT_USER,    IBOX_PT_PASS,    "PT iBox"),
    }

    raw: dict = {}

    def _fetch_vc(env, host, user, pwd):
        if not host:
            return f"vc_{env}", {"error": "not configured"}
        try:
            c = VCenterClient(host=host, username=user, password=pwd)
            return f"vc_{env}", c.get_summary()
        except Exception as e:
            return f"vc_{env}", {"error": "connection error"}

    def _fetch_srm(env, host, user, pwd):
        if not host:
            return f"srm_{env}", {"installed": False, "error": "not configured"}
        try:
            c = VCenterClient(host=host, username=user, password=pwd)
            return f"srm_{env}", c.get_srm_info()
        except Exception as e:
            return f"srm_{env}", {"installed": False, "error": "connection error"}

    def _fetch_ib(key, host, user, pwd, label):
        try:
            c = InfinidatClient(host=host, username=user, password=pwd)
            reps = c.get_replicas_detailed()
            return f"ib_{key}", {"label": label, "host": host, "replicas": reps, "error": None}
        except Exception as e:
            return f"ib_{key}", {"label": label, "host": host, "replicas": [], "error": "connection error"}

    tasks = []
    for env, (h, u, p) in vc_configs.items():
        tasks.append(lambda e=env, h=h, u=u, p=p: _fetch_vc(e, h, u, p))
        tasks.append(lambda e=env, h=h, u=u, p=p: _fetch_srm(e, h, u, p))
    for key, (h, u, p, lbl) in ib_configs.items():
        tasks.append(lambda k=key, h=h, u=u, p=p, l=lbl: _fetch_ib(k, h, u, p, l))

    with ThreadPoolExecutor(max_workers=9) as pool:
        for future in as_completed([pool.submit(t) for t in tasks]):
            try:
                k, v = future.result()
                raw[k] = v
            except Exception:
                pass

    # ── Build system status summary ──────────────────────────────────────
    systems = {}
    for env in ("tc", "pt", "fattal"):
        vc = raw.get(f"vc_{env}", {})
        err = vc.get("error")
        systems[f"vcenter_{env}"] = {
            "label":      f"vCenter {env.upper()}",
            "status":     "error" if err else "ok",
            "error":      err,
            "total_vms":  vc.get("total_vms", 0),
            "powered_on": vc.get("powered_on", 0),
            "dr_vms":     sum(1 for v in vc.get("vms", []) if v.get("is_dr_replica")),
        }
    for key, cfg in ib_configs.items():
        ib = raw.get(f"ib_{key}", {})
        reps = ib.get("replicas", [])
        active  = sum(1 for r in reps if r.get("state") == "ACTIVE")
        error   = sum(1 for r in reps if r.get("state") == "ERROR")
        susp    = sum(1 for r in reps if r.get("state") == "SUSPENDED")
        systems[f"ibox_{key}"] = {
            "label":           cfg[3],
            "host":            cfg[0],
            "status":          "error" if ib.get("error") else "ok",
            "error":           ib.get("error"),
            "replicas_total":  len(reps),
            "replicas_active": active,
            "replicas_error":  error,
            "replicas_susp":   susp,
        }

    # ── Build per-customer DR map from vCenter data ──────────────────────
    customer_map: dict = {}
    env_labels = {"tc": "TC", "pt": "PT", "fattal": "Fattal"}

    # Pass 1: build vm_name → sap_number map from production (non-DR) VMs.
    # SRM placeholder VMs have the same name as their production counterpart
    # but don't carry the SAP_NUMBER custom attribute. We resolve by name.
    vm_name_sap: dict = {}   # normalized_name → sap_number
    for env in ("tc", "pt", "fattal"):
        vc = raw.get(f"vc_{env}", {})
        if vc.get("error"):
            continue
        for vm in vc.get("vms", []):
            sap  = vm.get("sap_number", "")
            name = (vm.get("name") or "").strip().lower()
            if sap and name and not vm.get("is_dr_replica"):
                vm_name_sap[name] = sap

    # Pass 2: build customer DR map
    for env in ("tc", "pt", "fattal"):
        vc = raw.get(f"vc_{env}", {})
        if vc.get("error"):
            continue
        env_label = env_labels[env]
        for vm in vc.get("vms", []):
            sap = vm.get("sap_number", "")
            # For DR placeholder VMs without a sap_number, resolve via name lookup
            if not sap and vm.get("is_dr_replica"):
                sap = vm_name_sap.get((vm.get("name") or "").strip().lower(), "")
            if not sap:
                continue
            if sap not in customer_map:
                customer_map[sap] = {
                    "sap_number":    sap,
                    "customer_name": cust_map.get(sap, sap),
                    "prod_vms":      [],
                    "dr_vms":        [],
                    "envs":          {},
                }
            c = customer_map[sap]
            if vm.get("is_dr_replica"):
                c["dr_vms"].append({**vm, "_env": env_label})
                e = c["envs"].setdefault(env_label, {"prod": 0, "dr": 0})
                e["dr"] += 1
            else:
                c["prod_vms"].append({**vm, "_env": env_label})
                e = c["envs"].setdefault(env_label, {"prod": 0, "dr": 0})
                e["prod"] += 1

    # Filter: only customers that actually have DR VMs
    dr_customers = []
    for sap, c in customer_map.items():
        if not c["dr_vms"]:
            continue
        # Compute overall status
        dr_on  = sum(1 for v in c["dr_vms"] if v.get("power_state") == "poweredOn")
        dr_off = len(c["dr_vms"]) - dr_on
        if dr_on > 0:
            status = "active_failover"  # DR VMs are running — unusual
        else:
            status = "standby"
        dr_customers.append({
            "sap_number":    sap,
            "customer_name": c["customer_name"],
            "prod_count":    len(c["prod_vms"]),
            "dr_count":      len(c["dr_vms"]),
            "dr_on":         dr_on,
            "dr_off":        dr_off,
            "envs":          c["envs"],
            "status":        status,
            "prod_vms": [{
                "name":        v["name"],
                "power_state": v["power_state"],
                "num_cpu":     v.get("num_cpu", 0),
                "memory_gb":   v.get("memory_gb", 0),
                "storage_committed_gb": v.get("storage_committed_gb", 0),
                "host":        v.get("host", ""),
                "cluster":     v.get("cluster", ""),
                "_env":        v.get("_env", ""),
            } for v in c["prod_vms"][:50]],
            "dr_vms": [{
                "name":        v["name"],
                "power_state": v["power_state"],
                "cluster":     v.get("cluster", ""),
                "_env":        v.get("_env", ""),
            } for v in c["dr_vms"][:50]],
        })
    dr_customers.sort(key=lambda x: x["customer_name"])

    # ── Replicas per iBox ────────────────────────────────────────────────
    replicas = {
        "tc03":   raw.get("ib_tc03",   {}).get("replicas", []),
        "tc_old": raw.get("ib_tc_old", {}).get("replicas", []),
        "pt":     raw.get("ib_pt",     {}).get("replicas", []),
    }

    # ── SRM data ─────────────────────────────────────────────────────────
    srm = {
        "tc":     raw.get("srm_tc",     {"installed": False}),
        "pt":     raw.get("srm_pt",     {"installed": False}),
        "fattal": raw.get("srm_fattal", {"installed": False}),
    }

    _dr_result = {
        "updated_at": int(time.time()),
        "systems":    systems,
        "customers":  dr_customers,
        "replicas":   replicas,
        "srm":        srm,
    }
    # Cache only a clean result — never freeze a degraded/error aggregate for the full TTL
    if not any(s.get("status") == "error" for s in systems.values()):
        _dr_cache["summary"] = {"data": _dr_result, "ts": time.time()}
    return _dr_result


# ------------------------------------------------------------------ #
#  Device registry                                                     #
# ------------------------------------------------------------------ #

@app.get("/api/devices")
def api_list_devices(farm: Optional[str] = None, _: str = Depends(require_admin)):
    return list_devices(farm)

@app.post("/api/devices", status_code=201)
def api_create_device(d: DeviceIn, _: str = Depends(require_admin)):
    return create_device(
        farm=d.farm, name=d.name, dtype=d.type,
        role=d.role, host=d.host,
        ip_address=d.ip_address, ilo_address=d.ilo_address,
        username=d.username, password=d.password,
        port=d.port, protocol=d.protocol,
        notes=d.notes, enabled=d.enabled,
    )

@app.put("/api/devices/{device_id}")
def api_update_device(device_id: int, d: DeviceIn, _: str = Depends(require_admin)):
    result = update_device(
        device_id,
        farm=d.farm, name=d.name, type=d.type,
        role=d.role, host=d.host,
        ip_address=d.ip_address, ilo_address=d.ilo_address,
        username=d.username, password=d.password,
        port=d.port, protocol=d.protocol,
        notes=d.notes, enabled=1 if d.enabled else 0,
    )
    if not result:
        raise HTTPException(status_code=404, detail="Device not found")
    return result

@app.delete("/api/devices/{device_id}")
def api_delete_device(device_id: int, _: str = Depends(require_admin)):
    if not delete_device(device_id):
        raise HTTPException(status_code=404, detail="Device not found")
    return {"status": "deleted", "id": device_id}


# ------------------------------------------------------------------ #
#  Customers API                                                        #
# ------------------------------------------------------------------ #

@app.get("/api/customers")
def api_list_customers():
    return list_customers()

@app.get("/api/customers/map")
def api_customers_map():
    return get_customers_map()

@app.post("/api/customers", status_code=201)
def api_create_customer(c: CustomerIn, _: str = Depends(require_admin)):
    try:
        return create_customer(c.sap_number, c.customer_name)
    except Exception as e:
        print(f"[customers] create error: {e}")
        raise HTTPException(status_code=400, detail="Customer creation failed — check for duplicate SAP number")

@app.put("/api/customers/{customer_id}")
def api_update_customer(customer_id: int, c: CustomerIn, _: str = Depends(require_admin)):
    result = update_customer(customer_id, c.sap_number, c.customer_name)
    if not result:
        raise HTTPException(status_code=404, detail="Customer not found")
    return result

@app.delete("/api/customers/{customer_id}")
def api_delete_customer(customer_id: int, _: str = Depends(require_admin)):
    if not delete_customer(customer_id):
        raise HTTPException(status_code=404, detail="Customer not found")
    return {"status": "deleted", "id": customer_id}

_MAX_UPLOAD_BYTES = 5 * 1024 * 1024  # 5 MB

@app.post("/api/customers/import")
async def api_import_customers(file: UploadFile = File(...), _: str = Depends(require_admin)):
    fname = (file.filename or "").lower()
    if not (fname.endswith(".xlsx") or fname.endswith(".xls")):
        raise HTTPException(status_code=400, detail="Only .xlsx or .xls files are accepted")
    try:
        raw = await file.read(_MAX_UPLOAD_BYTES + 1)
        if len(raw) > _MAX_UPLOAD_BYTES:
            raise HTTPException(status_code=400, detail="File too large (max 5 MB)")
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        headers = [str(cell.value or "").strip() for cell in ws[1]]
        try:
            sap_col  = headers.index("מס. לקוח")
            name_col = headers.index("שם לקוח")
        except ValueError:
            sap_col, name_col = 0, 1
        pairs = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if len(row) <= max(sap_col, name_col):
                continue
            sap  = row[sap_col]
            name = row[name_col]
            if sap and name:
                pairs.append((str(sap).strip(), str(name).strip()))
        count = upsert_customers_bulk(pairs)
        return {"status": "ok", "imported": count}
    except Exception as e:
        raise HTTPException(status_code=400, detail="Import failed: check file format and content")


# ------------------------------------------------------------------ #
#  AI chatbot (internal Ollama server)                                 #
# ------------------------------------------------------------------ #

_AI_SYSTEM_PROMPT = (
    "You are the assistant embedded in an internal UCS/vCenter/storage/backup "
    "dashboard. Answer questions about hardware, capacity, VMs and infrastructure "
    "resources, and give practical recommendations based only on the data the user "
    "or the dashboard gives you in this conversation. If you don't have enough data "
    "to answer precisely, say so instead of guessing. Reply in the same language the "
    "user writes in (Hebrew or English). When the data contains a section 'STRUCTURED SERVER DATA' "
    "(metrics of one VM plus findings from the dashboard's rule engine), base your analysis on those "
    "numbers only, never invent values, and prioritise actions. Dashboard thresholds: CPU Ready per vCPU "
    "<5% ok, 5-10% warning, >=10% critical, judged over time not by one peak; Co-Stop >=3% means too many "
    "vCPUs; any memory swap or balloon means the ESXi host is squeezing the VM; disk latency <20 ms ok, "
    "20-30 ms slow, >=30 ms bad; ESXi vCPU:core ratio above 4:1 is a risk; high CPU with low Ready and "
    "Co-Stop means the VM itself is using the CPU (find the process), while high Ready means host "
    "contention (move the VM, do not add vCPUs). Database servers keep memory as cache on purpose. "
    "Do not overstate: judge severity from the findings, not from a single peak. If the verdict says no "
    "resource problem was found, answer that no action is needed on the hypervisor side and stop: do not "
    "invent a cause, and do not list generic checks unless the user says users are reporting slowness."
)

@app.post("/api/ai/chat")
def api_ai_chat(payload: AiChatIn, _: str = Depends(get_current_user)):
    system_content = _AI_SYSTEM_PROMPT
    if payload.context.strip():
        system_content += "\n\nCurrent dashboard data provided by the user:\n" + payload.context.strip()

    messages = [{"role": "system", "content": system_content}]
    messages += [{"role": m.role, "content": m.content} for m in payload.history[-20:]]
    messages.append({"role": "user", "content": payload.message})

    try:
        reply = ollama_client.chat(messages)
    except Exception as e:
        print(f"[ai] chat error: {e}")
        raise HTTPException(status_code=502, detail="AI server unavailable")
    return {"reply": reply}


@app.get("/api/ai/models", dependencies=[Depends(require_admin)])
def api_ai_models():
    """Admin: which model the assistant uses, what else is installed, and what is loaded (GPU/CPU)."""
    out = {"host": f"{ollama_client.OLLAMA_HOST}:{ollama_client.OLLAMA_PORT}"}
    for key, fn in (("configured", ollama_client.configured_model), ("installed", ollama_client.list_models),
                    ("running", ollama_client.running_models)):
        try:
            out[key] = fn()
        except Exception as e:
            out[key] = None
            out[key + "_error"] = f"{type(e).__name__}: {str(e)[:160]}"
    return out


@app.get("/api/ai/benchmark/start", dependencies=[Depends(require_admin)])
def api_ai_benchmark_start(model: str = ""):
    """Admin: run the 4-scenario benchmark on a model (default: the configured one) in the background."""
    import ai_benchmark
    try:
        names = [m.strip() for m in model.split(",") if m.strip()] or [ollama_client.configured_model()]
    except Exception:
        raise HTTPException(status_code=502, detail="AI server unavailable")
    return ai_benchmark.start(names, _AI_SYSTEM_PROMPT)


@app.get("/api/vcenter/perf/diagnose/batch/start", dependencies=[Depends(require_admin)])
def api_diag_batch_start(range: str = "7d", min_level: int = 3, ai: int = 0, limit: int = 80):
    """Admin: diagnose all VMs at/above a load level (3 = problematic) in the background and write
    logs/diagnose_batch_*.txt (optionally with the Ollama explanation of each) for review and calibration."""
    import diag_batch
    since, table = _perf_since(range)
    return diag_batch.start(range, since, table, max(1, min(min_level, 3)), bool(ai), max(1, min(limit, 300)), _AI_SYSTEM_PROMPT)


@app.get("/api/vcenter/perf/diagnose/batch/status", dependencies=[Depends(require_admin)])
def api_diag_batch_status():
    import diag_batch
    return diag_batch.status()


@app.get("/api/ai/benchmark/results", dependencies=[Depends(require_admin)])
def api_ai_benchmark_results(format: str = "json"):
    import ai_benchmark
    from fastapi.responses import PlainTextResponse
    data = ai_benchmark.results()
    if format == "text":
        return PlainTextResponse(ai_benchmark.as_text(data), media_type="text/plain; charset=utf-8")
    return data


# ------------------------------------------------------------------ #
#  Serve frontend static files                                         #
# ------------------------------------------------------------------ #

@app.get("/health")
def health_check():
    return {"status": "healthy"}


if os.path.isdir(FRONTEND_DIR):
    app.mount("/static", StaticFiles(directory=FRONTEND_DIR), name="static")

    @app.get("/")
    def serve_tc():
        return FileResponse(os.path.join(FRONTEND_DIR, "index.html"))

    @app.get("/pt")
    def serve_pt():
        return FileResponse(os.path.join(FRONTEND_DIR, "pt.html"))

    @app.get("/fattal")
    def serve_fattal():
        return FileResponse(os.path.join(FRONTEND_DIR, "fattal.html"))

    @app.get("/compare")
    def serve_compare():
        return FileResponse(os.path.join(FRONTEND_DIR, "compare.html"))

    @app.get("/devices")
    def serve_devices():
        return FileResponse(os.path.join(FRONTEND_DIR, "devices.html"))

    @app.get("/infinidat/tc")
    def serve_infinidat_tc():
        return FileResponse(os.path.join(FRONTEND_DIR, "infinidat_tc.html"))

    @app.get("/infinidat/tc-old")
    def serve_infinidat_tc_old():
        return FileResponse(os.path.join(FRONTEND_DIR, "infinidat_tc_old.html"))

    @app.get("/infinidat/pt")
    def serve_infinidat_pt():
        return FileResponse(os.path.join(FRONTEND_DIR, "infinidat_pt.html"))

    @app.get("/vcenter/pt")
    def serve_vcenter_pt():
        return FileResponse(os.path.join(FRONTEND_DIR, "vcenter_pt.html"))

    @app.get("/vcenter/tc")
    def serve_vcenter_tc():
        return FileResponse(os.path.join(FRONTEND_DIR, "vcenter_tc.html"))

    @app.get("/vcenter/tc8")
    def serve_vcenter_tc8():
        return FileResponse(os.path.join(FRONTEND_DIR, "vcenter_tc8.html"))

    @app.get("/vcenter/fattal")
    def serve_vcenter_fattal():
        return FileResponse(os.path.join(FRONTEND_DIR, "vcenter_fattal.html"))


    @app.get("/vcenter/cluster-detail")
    def serve_vcenter_cluster_detail():
        return FileResponse(os.path.join(FRONTEND_DIR, "vcenter_cluster_detail.html"))

    @app.get("/vcenter/changes")
    def serve_vcenter_changes():
        return FileResponse(os.path.join(FRONTEND_DIR, "vcenter_changes.html"))

    @app.get("/vcenter/overview")
    def serve_vcenter_overview():
        return FileResponse(os.path.join(FRONTEND_DIR, "vcenter_overview.html"))

    @app.get("/vcenter/vm-create")
    def serve_vcenter_vm_create():
        return FileResponse(os.path.join(FRONTEND_DIR, "vcenter_vm_create.html"))

    @app.get("/vcenter/history")
    def serve_vcenter_history():
        return FileResponse(os.path.join(FRONTEND_DIR, "vcenter_history.html"))

    @app.get("/vcenter/perf")
    def serve_vcenter_perf():
        return FileResponse(os.path.join(FRONTEND_DIR, "vcenter_perf.html"))

    @app.get("/vcenter/rightsizing")
    def serve_vcenter_rightsizing():
        return FileResponse(os.path.join(FRONTEND_DIR, "vcenter_rightsizing.html"))

    @app.get("/vcenter/reports")
    def serve_vcenter_reports():
        return FileResponse(os.path.join(FRONTEND_DIR, "vcenter_reports.html"))

    @app.get("/customers")
    def serve_customers_short():
        return FileResponse(os.path.join(FRONTEND_DIR, "customers.html"))

    @app.get("/management/customers")
    def serve_customers():
        return FileResponse(os.path.join(FRONTEND_DIR, "customers.html"))


    @app.get("/login")
    def serve_login():
        return FileResponse(os.path.join(FRONTEND_DIR, "login.html"))
    @app.get("/dr")
    def serve_dr():
        return FileResponse(os.path.join(FRONTEND_DIR, "dr.html"))

    @app.get("/backup")
    def serve_backup():
        return FileResponse(os.path.join(FRONTEND_DIR, "backup.html"))

    @app.get("/backup/pt")
    def serve_backup_pt():
        return FileResponse(os.path.join(FRONTEND_DIR, "backup_pt.html"))

    @app.get("/backup/lifesaver")
    def serve_backup_lifesaver():
        return FileResponse(os.path.join(FRONTEND_DIR, "backup_lifesaver.html"))

    @app.get("/backup/tc")
    def serve_backup_tc():
        return FileResponse(os.path.join(FRONTEND_DIR, "backup_tc.html"))

    @app.get("/backup/fattal")
    def serve_backup_fattal():
        return FileResponse(os.path.join(FRONTEND_DIR, "backup_fattal.html"))

    @app.get("/admin/overview")
    def serve_admin_overview():
        return FileResponse(os.path.join(FRONTEND_DIR, "admin_overview.html"))

    @app.get("/admin/permissions")
    def serve_admin_permissions(_: str = Depends(require_permissions_manager)):
        return FileResponse(os.path.join(FRONTEND_DIR, "permissions.html"))

    @app.get("/admin/billing")
    def serve_admin_billing(_: str = Depends(require_admin)):
        return FileResponse(os.path.join(FRONTEND_DIR, "billing_config.html"))


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("app:app", host="0.0.0.0", port=9000, reload=True)
