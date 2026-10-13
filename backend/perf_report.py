"""
Word (.docx) report of one VM for the vCenter Performance page: everything the screen shows for a server -
summary metrics, the diagnosis, the ESXi card, CPU topology, disks, the chart (a PNG drawn by the browser),
vCenter events and backup points - in Hebrew (right-to-left) or English.

Input is the "bundle" the app builds for the diagnosis (summary row, live host, disks, events, backups, vm_cfg,
the diagnosis result) plus the chart image; output is the .docx as bytes. python-docx only.
"""

import io
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import perf_diagnose as pd

_SEV_COLOR = {"crit": RGBColor(0xCF, 0x22, 0x2E), "warn": RGBColor(0x9A, 0x67, 0x00),
              "info": RGBColor(0x09, 0x69, 0xDA), "ok": RGBColor(0x1A, 0x7F, 0x37)}
_SEV_TXT = {"he": {"crit": "קריטי", "warn": "אזהרה", "info": "מידע", "ok": "תקין"},
            "en": {"crit": "Critical", "warn": "Warning", "info": "Info", "ok": "OK"}}

T = {
    "he": {
        "title": "דוח ביצועים", "env": "סביבה", "range": "טווח", "generated": "הופק", "by": "על ידי",
        "summary": "סיכום מדדים", "diag": "אבחון", "esxi": "ESXi (נתון חי)", "topo": "טופולוגיית CPU", "disks": "דיסקים (חי, ממוצע 5 דקות אחרונות)",
        "chart": "גרף", "events": "אירועי שרת מ-vCenter", "backups": "גיבויים וסנאפשוטים",
        "metric": "מדד", "value": "ערך", "vm": "שרת", "vcpu_ram": "vCPU / RAM (GB)",
        "cpu": "CPU", "ram": "RAM", "peak": "פיק", "avg": "ממוצע", "at": "ב-", "pct_time": "% מהזמן מעל 85%", "meter": "מד עומס",
        "ready": "CPU Ready", "costop": "Co-Stop", "limit": "CPU Limit", "mempress": "לחץ זיכרון",
        "limit_min": "דקות חסימה", "none": "אין", "verdict": "פסק", "cpu_answer": "האם צריך עוד מעבדים?", "todo": "מה לעשות",
        "host": "שרת פיזי", "state": "מצב", "load": "עומס", "vms": "שרתים", "ratio": "יחס vCPU:ליבה", "top": "הצרכנים הגדולים",
        "cur": "סידור נוכחי (סוקטים × ליבות)", "rec": "סידור מומלץ", "hotadd": "CPU Hot Add", "on": "מופעל", "off": "כבוי",
        "disk": "דיסק", "datastore": "Datastore", "size": "גודל", "read": "Read", "write": "Write", "iops": "IOPS (R/W)", "mbps": "MB/s (R/W)",
        "time": "זמן", "type": "סוג", "user": "מי ביצע", "server": "שרת גיבוי", "job": "Job", "kind": "סוג נקודה",
        "backup": "גיבוי", "replica": "רפליקציה / סנאפשוט", "no_events": "לא נמצאו אירועים בטווח.", "no_backups": "לא נמצאו נקודות גיבוי בטווח.",
        "no_chart": "(הגרף לא זמין)", "no_data": "אין נתונים בטווח שנבחר.",
        "levels": {3: "בעייתי", 2: "לתשומת לב", 1: "פיק חד-פעמי", 0: "תקין", -1: "מעט נתונים"},
        "footer": "הדוח הופק אוטומטית מדשבורד UCS. הספים והמסקנות הם כללי אצבע ואינם מחליפים בדיקה בתוך השרת. אירועי vCenter מגיעים מהיסטוריה של עד 30 יום.",
        "range_names": {"24h": "24 שעות אחרונות", "7d": "7 ימים", "30d": "30 יום", "90d": "90 יום", "180d": "6 חודשים"},
        "peak_hint": "(זמן ה-peak המדויק בסוגריים)",
    },
    "en": {
        "title": "Performance report", "env": "Environment", "range": "Range", "generated": "Generated", "by": "by",
        "summary": "Metrics summary", "diag": "Diagnosis", "esxi": "ESXi (live)", "topo": "CPU topology", "disks": "Disks (live, last 5 minutes average)",
        "chart": "Chart", "events": "Server events from vCenter", "backups": "Backups and snapshots",
        "metric": "Metric", "value": "Value", "vm": "Server", "vcpu_ram": "vCPU / RAM (GB)",
        "cpu": "CPU", "ram": "RAM", "peak": "Peak", "avg": "Average", "at": "at ", "pct_time": "% of time above 85%", "meter": "Load meter",
        "ready": "CPU Ready", "costop": "Co-Stop", "limit": "CPU Limit", "mempress": "Memory pressure",
        "limit_min": "minutes held back", "none": "none", "verdict": "Verdict", "cpu_answer": "Does it need more CPUs?", "todo": "What to do",
        "host": "Host", "state": "State", "load": "Load", "vms": "VMs", "ratio": "vCPU:core ratio", "top": "Top consumers",
        "cur": "Current layout (sockets x cores)", "rec": "Recommended layout", "hotadd": "CPU Hot Add", "on": "on", "off": "off",
        "disk": "Disk", "datastore": "Datastore", "size": "Size", "read": "Read", "write": "Write", "iops": "IOPS (R/W)", "mbps": "MB/s (R/W)",
        "time": "Time", "type": "Type", "user": "By", "server": "Backup server", "job": "Job", "kind": "Point type",
        "backup": "backup", "replica": "replica / snapshot", "no_events": "No events in the range.", "no_backups": "No backup points in the range.",
        "no_chart": "(chart not available)", "no_data": "No data in the selected range.",
        "levels": {3: "problematic", 2: "attention", 1: "one-off peak", 0: "ok", -1: "little data"},
        "footer": "Generated automatically by the UCS dashboard. Thresholds and conclusions are rules of thumb and do not replace a check inside the server. vCenter events cover up to 30 days.",
        "range_names": {"24h": "last 24 hours", "7d": "7 days", "30d": "30 days", "90d": "90 days", "180d": "6 months"},
        "peak_hint": "(exact peak time in brackets)",
    },
}


def _rtl_par(p):
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn("w:bidi")) is None:
        bidi = OxmlElement("w:bidi")
        pPr.insert_element_before(bidi, "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
                                  "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
                                  "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")


def _rtl_run(run):
    rPr = run._r.get_or_add_rPr()
    if rPr.find(qn("w:rtl")) is None:
        rtl = OxmlElement("w:rtl")
        rPr.insert_element_before(rtl, "w:cs", "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath")


def _table_rtl(table):
    tblPr = table._tbl.tblPr
    if tblPr.find(qn("w:bidiVisual")) is None:
        el = OxmlElement("w:bidiVisual")
        tblPr.insert_element_before(el, "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd", "w:tblLayout",
                                    "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription")


class _Doc:
    def __init__(self, lang: str):
        self.lang = lang if lang in ("he", "en") else "he"
        self.rtl = self.lang == "he"
        self.t = T[self.lang]
        self.d = Document()
        st = self.d.styles["Normal"]
        st.font.name = "Calibri"
        st.font.size = Pt(10)
        rpr = st.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        fonts.set(qn("w:cs"), "Arial")
        for sec in self.d.sections:
            sec.left_margin = sec.right_margin = Inches(0.8)
            sec.top_margin = sec.bottom_margin = Inches(0.7)

    # ---- text helpers -------------------------------------------------------------------------
    def par(self, container, text="", bold=False, color=None, size=None, italic=False, style=None, after=2):
        p = container.add_paragraph(style=style) if style else container.add_paragraph()
        if text != "":
            self.run(p, text, bold, color, size, italic)
        p.paragraph_format.space_after = Pt(after)
        if self.rtl:
            _rtl_par(p)
        return p

    def run(self, p, text, bold=False, color=None, size=None, italic=False):
        r = p.add_run(str(text))
        r.bold = bold
        r.italic = italic
        if color is not None:
            r.font.color.rgb = color
        if size:
            r.font.size = Pt(size)
        if self.rtl:
            _rtl_run(r)
        return r

    def heading(self, text):
        p = self.d.add_heading(level=2)
        self.run(p, text, bold=True, color=RGBColor(0x1F, 0x6F, 0xEB), size=13)
        if self.rtl:
            _rtl_par(p)
        return p

    def table(self, header, rows, widths=None):
        tb = self.d.add_table(rows=1, cols=len(header))
        tb.style = "Light Grid Accent 1"
        if self.rtl:
            _table_rtl(tb)
        for i, h in enumerate(header):
            c = tb.rows[0].cells[i]
            c.text = ""
            self.run(c.paragraphs[0], h, bold=True)
            if self.rtl:
                _rtl_par(c.paragraphs[0])
        for row in rows:
            cells = tb.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                color = None
                if isinstance(v, tuple):
                    v, color = v
                self.run(cells[i].paragraphs[0], "" if v is None else v, color=color)
                if self.rtl:
                    _rtl_par(cells[i].paragraphs[0])
        if widths:
            for row in tb.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Inches(w)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)
        return tb


def _n(v, d=1):
    return "-" if v is None else (("%." + str(d) + "f") % float(v)).rstrip("0").rstrip(".") if d else str(int(round(float(v))))


def _topology(vcpu, cfg, host):
    """(current, recommended, hot_add) like the page: physical sockets assumed = 2 per host."""
    if not vcpu or not cfg or not cfg[0]:
        return None
    cps, hot = int(cfg[0]), bool(cfg[1])
    vs = vcpu // cps
    node = int((host.get("cores") or 0) // 2) if host and host.get("cores") else 0
    rec = ""
    if node and vcpu <= node:
        rec = "\u202a1 \u00d7 %s\u202c" % vcpu
    elif node:
        nodes = -(-vcpu // node)
        rec = "\u202a%s \u00d7 %s\u202c" % (nodes, -(-vcpu // nodes))
    return "\u202a%s \u00d7 %s\u202c" % (vs, cps), rec, hot


def build_docx(bundle: dict, lang: str, chart_png: bytes, user: str, range_key: str) -> bytes:
    D = _Doc(lang)
    t = D.t
    L = D.lang
    s = bundle.get("summary") or {}
    host = bundle.get("host")
    disks = bundle.get("disks") or {}
    events = bundle.get("events") or []
    backups = bundle.get("backups") or {}
    vm_cfg = bundle.get("vm_cfg")
    res = bundle["result"]
    env, name = bundle["env"], bundle["name"]

    # ---- title -----------------------------------------------------------------------------------
    p = D.par(D.d, "%s: %s" % (t["title"], name), bold=True, size=18, color=RGBColor(0x1F, 0x23, 0x28), after=2)
    D.par(D.d, "%s: %s   |   %s: %s   |   %s %s %s %s" % (t["env"], env, t["range"], t["range_names"].get(range_key, range_key),
                                                       t["generated"], datetime.now().strftime("%d/%m/%Y %H:%M"), t["by"], user or "-"),
          color=RGBColor(0x65, 0x6D, 0x76), after=6)

    # ---- summary metrics -------------------------------------------------------------------------
    D.heading(t["summary"])
    if not s or not s.get("samples"):
        D.par(D.d, t["no_data"])
    else:
        n = max(s["samples"], 1)
        cpu_lv = pd._load_level(n, s.get("cpu_n_mid") or 0, s.get("cpu_n_high") or 0, s.get("cpu_hot_buckets") or 0, s.get("buckets") or 0)
        mem_lv = pd._load_level(n, s.get("mem_n_mid") or 0, s.get("mem_n_high") or 0, s.get("mem_hot_buckets") or 0, s.get("buckets") or 0)
        rdy_lv = pd._ready_level(s)
        lv = t["levels"]
        rows = [
            (t["vm"], name),
            (t["vcpu_ram"], "%s / %s" % (s.get("num_cpu"), _n(s.get("mem_gb")))),
            ("%s %s" % (t["cpu"], t["peak"]), "\u202a%s%%\u202c  %s%s" % (_n(s.get("max_cpu_pct")), t["at"], s.get("max_cpu_ts") or "-")),
            ("%s %s" % (t["cpu"], t["avg"]), "\u202a%s%%\u202c   (%s: \u202a%s%%\u202c)" % (_n(s.get("avg_cpu_pct")), t["pct_time"], _n(100.0 * (s.get("cpu_n_high") or 0) / n))),
            ("%s - %s" % (t["cpu"], t["meter"]), lv.get(cpu_lv, "-")),
            ("%s %s" % (t["ram"], t["peak"]), "\u202a%s%%\u202c  %s%s" % (_n(s.get("max_mem_pct")), t["at"], s.get("max_mem_ts") or "-")),
            ("%s %s" % (t["ram"], t["avg"]), "\u202a%s%%\u202c   (%s: \u202a%s%%\u202c)" % (_n(s.get("avg_mem_pct")), t["pct_time"], _n(100.0 * (s.get("mem_n_high") or 0) / n))),
            ("%s - %s" % (t["ram"], t["meter"]), lv.get(mem_lv, "-")),
            ("%s %s" % (t["ready"], t["peak"]), "\u202a%s%%\u202c  %s%s" % (_n(s.get("ready_max_pct")), t["at"], s.get("ready_max_ts") or "-")),
            ("%s %s" % (t["ready"], t["avg"]), "\u202a%s%%\u202c" % _n(s.get("ready_avg_pct"), 2)),
            ("%s - %s" % (t["ready"], t["meter"]), lv.get(rdy_lv, "-")),
            (t["costop"], "\u202a%s%%\u202c  %s%s" % (_n(s.get("cs_max_pct")), t["at"], s.get("cs_max_ts") or "-")),
            (t["limit"], ("%s %s" % (s.get("lim_n"), t["limit_min"])) if (s.get("lim_n") or 0) > 0 else t["none"]),
            (t["mempress"], "swap %s MB, balloon %s MB, latency \u202a%s%%\u202c" % (_n(s.get("mp_swap_max_mb")), _n(s.get("mp_balloon_max_mb")), _n(s.get("mp_lat_max_pct"), 2))),
        ]
        D.table([t["metric"], t["value"]], rows, widths=[2.6, 4.2])

    # ---- diagnosis -------------------------------------------------------------------------------
    D.heading(t["diag"])
    sev_of = {"crit": "crit", "warn": "warn", "info": "info", "ok": "ok"}
    D.par(D.d, "%s: %s" % (t["verdict"], res["verdict"][L]), bold=True, color=_SEV_COLOR.get(sev_of.get(res.get("level"), "info")), after=3)
    if res.get("cpu_answer"):
        p = D.par(D.d, "", after=4)
        D.run(p, "%s " % t["cpu_answer"], bold=True)
        D.run(p, res["cpu_answer"][L])
    for f in res.get("findings", []):
        col = _SEV_COLOR.get(f["sev"], _SEV_COLOR["info"])
        p = D.par(D.d, "", after=1)
        D.run(p, "[%s] " % _SEV_TXT[L][f["sev"]], bold=True, color=col)
        D.run(p, "%s - %s" % (f["area"], f["title"][L]), bold=True)
        D.par(D.d, f["detail"][L], after=1)
        if f.get("action") and f["action"][L]:
            p = D.par(D.d, "", after=5)
            D.run(p, "%s: " % t["todo"], bold=True, color=RGBColor(0x09, 0x69, 0xDA))
            D.run(p, f["action"][L], color=RGBColor(0x09, 0x69, 0xDA))

    # ---- ESXi + topology -------------------------------------------------------------------------
    D.heading(t["esxi"])
    if host:
        D.table([t["metric"], t["value"]], [
            (t["host"], host.get("name")),
            (t["state"], ("%s%s" % (host.get("level"), (" - " + "; ".join(host.get("reasons_en" if L == "en" else "reasons") or [])) if host.get("reasons") else ""))),
            (t["load"], "CPU \u202a%s%%\u202c, RAM \u202a%s%%\u202c (%s GB), CPU Ready \u202a%s%%\u202c (max \u202a%s%%\u202c)" % (_n(host.get("cpu_pct")), _n(host.get("mem_pct")), _n(host.get("mem_total_gb")),
                                                                                  _n(host.get("ready_avg"), 2), _n(host.get("ready_max"), 2))),
            (t["vms"], "%s VMs, %s vCPU / %s cores  (%s)" % (host.get("vms"), host.get("vcpus"), host.get("cores"), "%s:1" % _n(host.get("vcpu_ratio")))),
            (t["top"], ", ".join("%s \u202a%s%%\u202c" % (x.get("name"), _n(x.get("cpu_pct"))) for x in (host.get("top_vms") or [])) or "-"),
        ], widths=[1.8, 5.0])
    else:
        D.par(D.d, "-")
    topo = _topology(s.get("num_cpu"), vm_cfg, host)
    if topo:
        D.par(D.d, t["topo"], bold=True, after=1)
        D.table([t["metric"], t["value"]], [
            (t["cur"], topo[0]), (t["rec"], topo[1] or "-"), (t["hotadd"], t["on"] if topo[2] else t["off"]),
        ], widths=[2.8, 4.0])

    # ---- disks -----------------------------------------------------------------------------------
    D.heading(t["disks"])
    dl = disks.get("disks") or []
    if dl:
        D.table([t["disk"], t["datastore"], t["size"], t["read"], t["write"], t["iops"], t["mbps"]], [
            (d.get("label"), d.get("datastore"), "%s GB" % _n(d.get("size_gb")),
             "%s ms" % _n(d.get("read_ms")), "%s ms" % _n(d.get("write_ms")),
             "%s / %s" % (_n(d.get("read_iops")), _n(d.get("write_iops"))), "%s / %s" % (_n(d.get("read_mbps"), 2), _n(d.get("write_mbps"), 2)))
            for d in dl], widths=[1.3, 1.9, 0.8, 0.7, 0.7, 0.9, 0.9])
    else:
        D.par(D.d, "-")

    # ---- chart -----------------------------------------------------------------------------------
    D.heading(t["chart"])
    if chart_png:
        try:
            D.d.add_picture(io.BytesIO(chart_png), width=Inches(6.8))
            D.par(D.d, ("CPU %% - %s | RAM %% - %s | %s" % ("כחול" if L == "he" else "blue", "סגול" if L == "he" else "purple", t["peak_hint"])),
                  color=RGBColor(0x65, 0x6D, 0x76), size=8)
        except Exception:
            D.par(D.d, t["no_chart"])
    else:
        D.par(D.d, t["no_chart"])

    # ---- events ----------------------------------------------------------------------------------
    D.heading(t["events"])
    ev = [e for e in events if e.get("kind") != "snapshot"][-15:]
    if ev:
        D.table([t["time"], t["type"], t["user"]], [(e.get("time"), e.get("label") if L == "he" else e.get("label_en", e.get("kind")), e.get("user") or "-") for e in reversed(ev)],
                widths=[1.5, 3.4, 1.9])
    else:
        D.par(D.d, t["no_events"])

    # ---- backups ---------------------------------------------------------------------------------
    D.heading(t["backups"])
    pts = backups.get("points") or []
    if pts:
        nb = len([p for p in pts if p.get("kind") == "backup"])
        nr = len(pts) - nb
        D.par(D.d, "%s: %s   |   %s: %s" % (t["backup"], nb, t["replica"], nr), bold=True, after=3)
        D.table([t["time"], t["server"], t["job"], t["kind"]], [
            (p.get("time"), p.get("server"), p.get("job"), t["backup"] if p.get("kind") == "backup" else t["replica"]) for p in pts[-15:][::-1]],
            widths=[1.5, 1.2, 2.9, 1.2])
    else:
        D.par(D.d, t["no_backups"])

    D.par(D.d, "")
    D.par(D.d, t["footer"], italic=True, size=8, color=RGBColor(0x65, 0x6D, 0x76))
    out = io.BytesIO()
    D.d.save(out)
    return out.getvalue()
